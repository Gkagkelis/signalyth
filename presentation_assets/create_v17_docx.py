from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('/mnt/data')
APP=OUT/'SIGNALYTH-app-v1.7'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(cell, color='D9D1C6'):
    tcPr=cell._tc.get_or_add_tcPr()
    tcBorders=tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders=OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge
        element=tcBorders.find(qn(tag))
        if element is None:
            element=OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn('w:val'),'single')
        element.set(qn('w:sz'),'4')
        element.set(qn('w:space'),'0')
        element.set(qn('w:color'),color)

def add_p(text, style=None, bold=False, size=None, color=None):
    p=doc.add_paragraph(style=style)
    r=p.add_run(text)
    r.bold=bold
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return p

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.55); sec.bottom_margin=Inches(0.55); sec.left_margin=Inches(0.62); sec.right_margin=Inches(0.62)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5)
for st in ['Title','Heading 1','Heading 2']:
    styles[st].font.name='Aptos Display'

p=doc.add_paragraph()
r=p.add_run('SIGNALYTH v1.7 - Internal QA Brief')
r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor(23,22,20)
add_p('STEP 8.3 - WOW Presentation + Decision Intelligence + Client Branding', size=11, color='6F6A62')
add_p('This build upgrades v1.6 from evidence-complete reporting to a client-facing presentation system designed to be visually impressive, decision-oriented and still evidence-safe.', size=10)

add_p('What changed', style='Heading 1')
items=[
 'Client Profiles: client logo upload, brand accent, default report language.',
 'Agency branding: your company logo stored once and shown discreetly on cover/closing, not on every slide.',
 'Powered by SIGNALYTH: kept subtle as technology credit.',
 'WOW Visual Intelligence: donuts, timelines, heatmaps, rankings, bubble/matrix views, map logic only when reliable geography exists.',
 'Event/Crisis/Opportunity detection: spikes are evaluated with guardrails before being labeled.',
 'Recommendations Engine: every recommendation must have a finding, evidence IDs, priority, timing and monitor metrics.',
 'Full mode keeps evidence-heavy appendix so comments/mentions are visible to the client.'
]
for it in items: add_p('• '+it)

add_p('Hard gates before client PPTX export', style='Heading 1')
gates=[
 ['Gate','Pass condition'],
 ['Visual reason','Every chart/visual has a data reason and presentation role.'],
 ['No fake drama','Crisis/Opportunity labels require threshold + independent voices + source spread + evidence.'],
 ['Recommendation evidence','No generic recommendations; every action must trace to findings and evidence.'],
 ['Client branding','Client logo leads; agency logo and SIGNALYTH remain discreet.'],
 ['Bilingual QA','EN/EL titles, labels and recommendation text are checked; original comments remain unchanged.'],
 ['Evidence density','Full mode includes long lists of comments/mentions, not only isolated examples.']
]
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
for i,h in enumerate(gates[0]):
    c=t.rows[0].cells[i]; c.text=h; shade(c,'EFE7DC'); borders(c)
    for p in c.paragraphs: p.runs[0].bold=True
for row in gates[1:]:
    cells=t.add_row().cells
    for i,v in enumerate(row):
        cells[i].text=v; borders(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

add_p('Benchmark outputs', style='Heading 1')
add_p('The build creates English and Greek benchmark decks with client/agency branding hierarchy, event detection slides, recommendation slides, WOW visualization examples, and evidence-heavy appendix tables.')

add_p('Remaining production boundary', style='Heading 1')
add_p('This is a controlled benchmark. The final proof remains a live end-to-end run using real Apify and OpenAI credentials, followed by QA of the resulting client PPTX, internal DOCX/PDF, and evidence exports.', bold=True, color='A9453F')

add_p('QA summary', style='Heading 1')
summary=[['Area','Status'],['PPTX EN/EL generated','PASS'],['PDF export','PASS'],['Slide render visual smoke test','PASS'],['Native editable charts present','PASS'],['Client/agency/SIGNALYTH branding hierarchy','PASS'],['Crisis guardrail contract','PASS'],['Evidence-linked recommendations contract','PASS'],['Clean ZIP re-test','PASS']]
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(summary[0]):
    c=t.rows[0].cells[i]; c.text=h; shade(c,'E3D6C5'); borders(c); c.paragraphs[0].runs[0].bold=True
for row in summary[1:]:
    cells=t.add_row().cells
    for i,v in enumerate(row):
        cells[i].text=v; borders(cells[i])

# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]
    footer.text='SIGNALYTH v1.7 · Internal QA Brief'
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in footer.runs: r.font.size=Pt(7); r.font.color.rgb=RGBColor(111,106,98)

out=OUT/'SIGNALYTH-Step8-WOW-Decision-Internal-Brief-v1.7.docx'
doc.save(out)
print(out)
