"""③ Continuation contract for the Top-comments slides.

Rule under test: a comment is NEVER truncated and never shrunk below the font
ladder — when the Top-N stops fitting, the slide continues as (1/2)-(2/2) with
the same design system, continued rank numbering, and whole-list chips.
"""
import unittest

from pptx import Presentation
from pptx.util import Inches

from app.services.presentation import (
    _top_comments_parts, _tc_fits, _tc_commentary_slice, _tc_card_heights,
    _render_top_comments, _TC_READABLE_MIN, _TC_FONT_LADDER,
)


def _entry(text, score, idx=1, author="user"):
    return {"author": author, "author_idx": idx, "platform": "", "date": "2/9",
            "score": score, "impact": 0.5, "text": text}


def short_entries(n=10):
    return [_entry(f"Σύντομο σχόλιο νούμερο {i} για το προϊόν.", round(0.9 - i * 0.05, 2), i + 1)
            for i in range(n)]


def long_entries(n=10, chars=420):
    base = ("Αναλυτική τοποθέτηση με πλήρη επιχειρηματολογία, παραδείγματα και συγκρίσεις "
            "που δεν πρέπει να χαθεί ούτε λέξη από αυτήν. ")
    return [_entry((base * 20)[:chars] + f" [#{i}]", round(0.9 - i * 0.05, 2), i + 1)
            for i in range(n)]


class SplitDecisionTests(unittest.TestCase):
    def test_short_top10_stays_on_one_slide(self):
        self.assertEqual(_top_comments_parts(short_entries()), [(0, 10)])

    def test_long_top10_splits_into_two_balanced_parts(self):
        entries = long_entries()
        parts = _top_comments_parts(entries)
        self.assertEqual(len(parts), 2, parts)
        # Contiguous full coverage, no overlap, no loss.
        self.assertEqual(parts[0][0], 0)
        self.assertEqual(parts[-1][1], len(entries))
        self.assertEqual(parts[0][1], parts[1][0])
        # Balanced by content height: neither page carries almost everything.
        heights = _tc_card_heights(entries, _TC_FONT_LADDER[0])
        h1 = sum(heights[parts[0][0]:parts[0][1]])
        h2 = sum(heights[parts[1][0]:parts[1][1]])
        self.assertLess(abs(h1 - h2) / max(h1, h2), 0.35)

    def test_every_part_fits_unscaled_on_the_ladder(self):
        entries = long_entries()
        for s, e in _top_comments_parts(entries):
            self.assertIsNotNone(_tc_fits(entries[s:e]))

    def test_pathological_volume_grows_parts_instead_of_truncating(self):
        entries = long_entries(n=10, chars=1900)  # near the 2000-char cap
        parts = _top_comments_parts(entries)
        self.assertGreaterEqual(len(parts), 2)
        self.assertEqual(parts[0][0], 0)
        self.assertEqual(parts[-1][1], len(entries))
        for (s1, e1), (s2, e2) in zip(parts, parts[1:]):
            self.assertEqual(e1, s2)
        for s, e in parts:
            self.assertIsNotNone(_tc_fits(entries[s:e]))

    def test_single_or_empty_list_never_splits(self):
        self.assertEqual(_top_comments_parts([]), [(0, 0)])
        self.assertEqual(_top_comments_parts(long_entries(1)), [(0, 1)])


class CommentaryDistributionTests(unittest.TestCase):
    PARAS = [{"lead": f"L{i}.", "body": f"B{i}"} for i in range(3)]

    def test_paragraphs_distributed_without_repetition(self):
        e = long_entries()
        p1 = _tc_commentary_slice(self.PARAS, 1, 2, e[:5], 0, 5, "el")
        p2 = _tc_commentary_slice(self.PARAS, 2, 2, e[5:], 5, 10, "el")
        self.assertEqual([x["lead"] for x in p1] + [x["lead"] for x in p2],
                         [x["lead"] for x in self.PARAS])

    def test_prose_front_loads_and_empty_slot_gets_continuation_note(self):
        e = long_entries()
        p1 = _tc_commentary_slice([self.PARAS[0]], 1, 2, e[:5], 0, 5, "el")
        self.assertEqual(p1, [self.PARAS[0]])  # the single paragraph opens page 1
        p2 = _tc_commentary_slice([self.PARAS[0]], 2, 2, e[5:], 5, 10, "el")
        self.assertEqual(len(p2), 1)
        self.assertIn("6–10", p2[0]["body"])
        p2en = _tc_commentary_slice([], 2, 2, e[5:], 5, 10, "en")
        self.assertIn("Ranks 6–10", p2en[0]["body"])


class ContinuationRenderTests(unittest.TestCase):
    def _render(self, entries, part, all_scores):
        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _render_top_comments(slide, entries, [{"lead": "Lead.", "body": "Body"}],
                             "positive", "el", {"client": "OPAP", "topic": "EJ"}, {},
                             part=part, all_scores=all_scores)
        return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]

    def test_second_page_continues_rank_numbering(self):
        entries = long_entries()
        texts = self._render(entries[5:], {"start": 5}, [e["score"] for e in entries])
        self.assertTrue({"6", "7", "8", "9", "10"} <= set(texts))
        for lone in ("1", "2", "3", "4", "5"):
            self.assertNotIn(lone, texts)  # no restarted badges on page 2

    def test_chips_show_whole_list_numbers_on_both_pages(self):
        entries = long_entries()
        scores = [e["score"] for e in entries]
        t1 = self._render(entries[:5], {"start": 0}, scores)
        t2 = self._render(entries[5:], {"start": 5}, scores)
        avg = sum(scores) / len(scores)
        chip = f"+{avg:.2f}"
        self.assertTrue(any(chip in t for t in t1))
        self.assertTrue(any(chip in t for t in t2))
        self.assertTrue(any("Top 10" in t for t in t1))
        self.assertTrue(any("Top 10" in t for t in t2))

    def test_no_truncation_every_word_survives_the_split(self):
        entries = long_entries()
        joined = " ".join(self._render(entries[:5], {"start": 0}, None)) + " " + \
                 " ".join(self._render(entries[5:], {"start": 5}, None))
        for e in entries:
            self.assertIn(e["text"], joined)


if __name__ == "__main__":
    unittest.main()
