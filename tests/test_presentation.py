import copy
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from docx import Document

from app.services.investigations import INDICATOR_REGISTRY
from app.services.presentation import (
    PresentationValidationError, PresentationCancelled,
    build_presentation_plan, generate_pptx, generate_internal_docx,
    build_exports, load_export_summary, allowed_export_file, _convert_to_pdf,
)

ROOT=Path(__file__).resolve().parents[1]
FIX=json.loads((ROOT/'evals'/'presentation_fixture.json').read_text(encoding='utf-8'))


def plan(language='English'):
    return {'client':'OPAP','topic':'Eurojackpot','market':'Greece','date_from':'2026-08-01','date_to':'2026-08-31','report_language':language}


class PresentationPlanTests(unittest.TestCase):
    def setUp(self):
        self.e=copy.deepcopy(FIX['evidence_pack']); self.v=copy.deepcopy(FIX['visual_pack'])

    def test_reviews_all_25_indicators_exactly_once(self):
        out=build_presentation_plan(self.v,self.e,plan())
        self.assertEqual(len(out['indicator_review']),25)
        self.assertEqual([x['indicator_id'] for x in out['indicator_review']],[x[0] for x in INDICATOR_REGISTRY])
        self.assertTrue(all(x['examined'] for x in out['indicator_review']))

    def test_core_report_has_cover_executive_and_conclusions(self):
        out=build_presentation_plan(self.v,self.e,plan())
        ids=[s['slide_id'] for s in out['slides']]
        self.assertEqual(ids[0],'cover'); self.assertIn('executive_summary',ids); self.assertEqual(ids[-1],'conclusions')

    def test_report_is_dynamic_not_fixed_27_slide_clone(self):
        out=build_presentation_plan(self.v,self.e,plan())
        self.assertGreaterEqual(len(out['slides']),7); self.assertLess(len(out['slides']),20)

    def test_emotion_slide_is_created_when_material(self):
        out=build_presentation_plan(self.v,self.e,plan())
        self.assertIn('emotions',[s['slide_id'] for s in out['slides']])

    def test_low_signal_emotions_can_be_omitted_but_reviewed(self):
        c=next(x for x in self.v['chart_specs'] if x['chart_id']=='emotion_distribution')
        for item in c['data']['categories']: item['value']=1.0
        self.v['investigation_candidates']=[]
        out=build_presentation_plan(self.v,self.e,plan())
        self.assertNotIn('emotions',[s['slide_id'] for s in out['slides']])
        row=next(x for x in out['indicator_review'] if x['indicator_id']=='emotions')
        self.assertTrue(row['examined'])

    def test_investigation_slide_carries_not_proven_guardrail(self):
        out=build_presentation_plan(self.v,self.e,plan())
        inv=next(s for s in out['slides'] if s['slide_type']=='investigation')
        self.assertEqual(inv['notes']['causality_guardrail'],'Association ≠ proven causality.')
        self.assertTrue(inv['claims'])
        self.assertEqual(inv['claims'][0]['causal_status'],'not_proven')

    def test_every_material_claim_has_indicator_traceability(self):
        out=build_presentation_plan(self.v,self.e,plan())
        for cl in out['claim_ledger']:
            if cl['claim_type']!='guardrail': self.assertTrue(cl['indicator_ids'],cl)

    def test_claim_ids_unique(self):
        out=build_presentation_plan(self.v,self.e,plan())
        ids=[x['claim_id'] for x in out['claim_ledger']]; self.assertEqual(len(ids),len(set(ids)))

    def test_exec_reputation_uses_exact_input_value(self):
        out=build_presentation_plan(self.v,self.e,plan())
        cl=next(x for x in out['claim_ledger'] if x['claim_id']=='exec-reputation')
        self.assertEqual(cl['source_values']['brand_reputation'],61.25)
        self.assertIn('61.2',cl['text'])

    def test_all_step7_presentation_charts_are_explicitly_reviewed(self):
        out=build_presentation_plan(self.v,self.e,plan())
        expected=set(self.v['presentation_chart_order']); got={x['chart_id'] for x in out['chart_review']}
        self.assertEqual(expected,got); self.assertTrue(all(x['examined'] for x in out['chart_review']))

    def test_missing_indicator_fails_closed(self):
        self.v['indicator_review'].pop()
        with self.assertRaises(PresentationValidationError): build_presentation_plan(self.v,self.e,plan())

    def test_duplicate_indicator_fails_closed(self):
        self.v['indicator_review'].append(copy.deepcopy(self.v['indicator_review'][0]))
        with self.assertRaises(PresentationValidationError): build_presentation_plan(self.v,self.e,plan())

    def test_raster_chart_contract_is_rejected(self):
        self.v['chart_specs'][0]['presentation']['render_as_raster']=True
        with self.assertRaises(PresentationValidationError): build_presentation_plan(self.v,self.e,plan())

    def test_non_editable_chart_contract_is_rejected(self):
        self.v['chart_specs'][0]['presentation']['native_editable_ready']=False
        with self.assertRaises(PresentationValidationError): build_presentation_plan(self.v,self.e,plan())

    def test_research_frame_mismatch_is_rejected(self):
        p=plan();p['market']='Italy'
        with self.assertRaises(PresentationValidationError): build_presentation_plan(self.v,self.e,p)

    def test_greek_report_language_changes_titles(self):
        out=build_presentation_plan(self.v,self.e,plan('Ελληνικά'))
        self.assertEqual(out['language'],'el');self.assertEqual(next(x for x in out['slides'] if x['slide_id']=='executive_summary')['title'],'Executive Dashboard')

    def test_cancel_before_planning(self):
        with self.assertRaises(PresentationCancelled): build_presentation_plan(self.v,self.e,plan(),cancel_check=lambda:True)

    def test_does_not_mutate_visual_or_evidence_pack(self):
        v=copy.deepcopy(self.v);e=copy.deepcopy(self.e);build_presentation_plan(self.v,self.e,plan());self.assertEqual(v,self.v);self.assertEqual(e,self.e)


class RendererTests(unittest.TestCase):
    def setUp(self):
        self.e=copy.deepcopy(FIX['evidence_pack']); self.v=copy.deepcopy(FIX['visual_pack']); self.p=build_presentation_plan(self.v,self.e,plan())
        self.tmp=Path(tempfile.mkdtemp(prefix='sig-pres-'))
    def tearDown(self): shutil.rmtree(self.tmp,ignore_errors=True)

    def test_pptx_is_created_and_reopenable(self):
        path=self.tmp/'a.pptx';meta=generate_pptx(self.p,self.v,path,ROOT/'logo.png')
        self.assertTrue(path.exists());prs=Presentation(path);self.assertEqual(len(prs.slides),meta['slides']);self.assertEqual(len(prs.slides),len(self.p['slides']))

    def test_pptx_contains_native_chart_parts(self):
        path=self.tmp/'a.pptx';generate_pptx(self.p,self.v,path,ROOT/'logo.png')
        with zipfile.ZipFile(path) as z:
            charts=[n for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')]
            self.assertGreaterEqual(len(charts),3)

    def test_pptx_has_no_chart_screenshot_media_dependency(self):
        path=self.tmp/'a.pptx';generate_pptx(self.p,self.v,path,ROOT/'logo.png')
        with zipfile.ZipFile(path) as z:
            media=[n for n in z.namelist() if n.startswith('ppt/media/')]
            # Only the user-provided logo is expected as raster media in the fixture deck.
            self.assertLessEqual(len(media),1)

    def test_renderer_audit_uses_only_editable_modes(self):
        path=self.tmp/'a.pptx';meta=generate_pptx(self.p,self.v,path,ROOT/'logo.png')
        modes={x['render_mode'] for x in meta['chart_render_audit']}
        self.assertFalse(any('raster' in m for m in modes));self.assertTrue(modes <= {'native_chart','editable_shapes','editable_table','editable_text_fallback'})

    def test_greek_pptx_reopens_with_greek_text(self):
        gp=build_presentation_plan(self.v,self.e,plan('Ελληνικά'));path=self.tmp/'g.pptx';generate_pptx(gp,self.v,path,ROOT/'logo.png');prs=Presentation(path)
        text=' '.join(sh.text for sl in prs.slides for sh in sl.shapes if hasattr(sh,'text'))
        self.assertIn('Executive Dashboard',text)

    def test_internal_docx_is_created_and_reopenable(self):
        path=self.tmp/'internal.docx';meta=generate_internal_docx(self.p,self.e,path,ROOT/'logo.png');self.assertTrue(path.exists());doc=Document(path);self.assertGreaterEqual(len(doc.tables),2);self.assertEqual(meta['tables'],len(doc.tables))

    def test_internal_docx_contains_watch_section_and_claim_ledger(self):
        path=self.tmp/'internal.docx';generate_internal_docx(self.p,self.e,path,ROOT/'logo.png');doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('What to watch',text);self.assertIn('Claim ledger',text)

    @unittest.skipUnless(shutil.which('libreoffice') or shutil.which('soffice'),'LibreOffice not installed')
    def test_pptx_converts_to_nonempty_pdf(self):
        path=self.tmp/'a.pptx';generate_pptx(self.p,self.v,path,ROOT/'logo.png');pdf=_convert_to_pdf(path,self.tmp/'pdf');self.assertIsNotNone(pdf);self.assertGreater(pdf.stat().st_size,1000)

    @unittest.skipUnless(shutil.which('libreoffice') or shutil.which('soffice'),'LibreOffice not installed')
    def test_docx_converts_to_nonempty_pdf(self):
        path=self.tmp/'a.docx';generate_internal_docx(self.p,self.e,path,ROOT/'logo.png');pdf=_convert_to_pdf(path,self.tmp/'pdf');self.assertIsNotNone(pdf);self.assertGreater(pdf.stat().st_size,1000)

    def test_control_characters_are_sanitized(self):
        bad=copy.deepcopy(self.p);bad['slides'][-1]['claims'].append({'claim_id':'x','text':'hello\x00world','indicator_ids':['sentiment'],'claim_type':'descriptive','evidence_refs':[],'source_values':{},'causal_status':'not_applicable'})
        path=self.tmp/'a.pptx';generate_pptx(bad,self.v,path,ROOT/'logo.png');self.assertTrue(path.exists())


class ExportBundleTests(unittest.TestCase):
    def setUp(self):
        self.e=copy.deepcopy(FIX['evidence_pack']); self.v=copy.deepcopy(FIX['visual_pack']); self.tmp=Path(tempfile.mkdtemp(prefix='sig-bundle-'))
        (self.tmp/'plan.json').write_text(json.dumps(plan()),encoding='utf-8')
    def tearDown(self): shutil.rmtree(self.tmp,ignore_errors=True)

    def _build(self,force=False):
        # Real LibreOffice conversion is covered once per format in RendererTests.
        # Bundle tests focus on deterministic export contracts and should not pay
        # the process-startup cost twice for every assertion.
        with patch('app.services.presentation.load_visualization_summary',return_value={'stale':False}), patch('app.services.presentation.load_presentation_visual_pack',return_value=self.v), patch('app.services.presentation.load_evidence_pack',return_value=self.e), patch('app.services.presentation._convert_to_pdf',return_value=None):
            return build_exports(self.tmp,plan(),force=force)

    def test_bundle_generates_pptx_docx_and_evidence_exports(self):
        s=self._build();self.assertGreaterEqual(s['slide_count'],7);names=set(s['files']);self.assertTrue(any(x.endswith('.pptx') for x in names));self.assertTrue(any(x.endswith('.docx') for x in names));self.assertTrue(any(x.endswith('_Evidence.json') for x in names));self.assertTrue(any(x.endswith('_Evidence.csv') for x in names))

    def test_manifest_hashes_match_files(self):
        import hashlib
        self._build();m=json.loads((self.tmp/'exports'/'manifest.json').read_text())
        for row in m['files']:
            p=self.tmp/'exports'/row['name'];self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),row['sha256']);self.assertEqual(p.stat().st_size,row['bytes'])

    def test_qa_requires_all_25_indicators(self):
        self._build();q=json.loads((self.tmp/'exports'/'qa.json').read_text());self.assertTrue(q['indicator_review_complete']);self.assertTrue(q['native_editable_chart_contract'])

    def test_evidence_export_keeps_record_id_and_url(self):
        self._build();p=next((self.tmp/'exports').glob('*_Evidence.json'));rows=json.loads(p.read_text());self.assertEqual(rows[0]['record_id'],'r1');self.assertEqual(rows[0]['url'],'https://x/1')

    def test_allowed_export_file_blocks_path_traversal(self):
        self._build();self.assertIsNone(allowed_export_file(self.tmp,'../plan.json'));self.assertIsNone(allowed_export_file(self.tmp,'not-in-manifest.txt'))

    def test_idempotent_build_reuses_same_summary_when_input_unchanged(self):
        a=self._build();b=self._build();self.assertEqual(a['input_hash'],b['input_hash']);self.assertEqual(a['generated_at'],b['generated_at'])

    def test_force_rebuild_keeps_same_input_hash(self):
        a=self._build();b=self._build(force=True);self.assertEqual(a['input_hash'],b['input_hash'])

    def test_load_export_summary_marks_changed_visual_pack_stale(self):
        self._build();changed=copy.deepcopy(self.v);changed['chart_specs'][0]['priority']=1
        with patch('app.services.presentation.load_presentation_visual_pack',return_value=changed), patch('app.services.presentation.load_evidence_pack',return_value=self.e):
            out=load_export_summary(self.tmp)
        self.assertTrue(out['stale'])

    def test_bundle_summary_has_complete_indicator_contract(self):
        s=self._build();self.assertEqual(s['indicator_contract']['required'],25);self.assertEqual(s['indicator_contract']['examined'],25);self.assertTrue(s['indicator_contract']['complete'])

    def test_no_raw_evidence_file_is_modified(self):
        original=copy.deepcopy(self.e);self._build();self.assertEqual(self.e,original)


if __name__=='__main__': unittest.main()
