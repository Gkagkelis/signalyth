"""④ Brand Reputation per channel on the Executive Dashboard.

Contract under test:
- channels come from the run's own source_comparison rows (per-platform indices
  produced by the SAME methodology transformation as the headline score);
- n = reputation-eligible records sits next to every score;
- platform names are visible ONLY in aggregate views (this strip); individual
  comment cards never carry a platform;
- (1/2)-(2/2) continuation pages share their part-1 section number.
"""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.services import intelligence as intel
from app.services.presentation import (
    _executive_dashboard_data, _render_executive_dashboard, _channel_label,
    _render_top_comments, generate_pptx, _upper_label,
)

ROOT = Path(__file__).resolve().parents[1]
FIX = json.loads((ROOT / "evals" / "presentation_fixture.json").read_text(encoding="utf-8"))


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _blank_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


class DashboardChannelDataTests(unittest.TestCase):
    def setUp(self):
        self.v = copy.deepcopy(FIX["visual_pack"])
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_channels_come_from_source_comparison_with_n(self):
        dash = _executive_dashboard_data(self.folder, self.v)
        self.assertIsNotNone(dash)
        chans = dash["channels"]
        src = next(c for c in self.v["chart_specs"] if c["chart_id"] == "source_comparison")
        rows = {r["source"]: r for r in src["data"]["rows"]}
        self.assertEqual({c["name"] for c in chans}, set(rows))  # only the run's platforms
        for c in chans:
            self.assertEqual(c["n"], rows[c["name"]]["reputation_eligible"])
            self.assertAlmostEqual(c["score"], round(float(rows[c["name"]]["brand_reputation"]), 1))
        self.assertEqual([c["n"] for c in chans], sorted((c["n"] for c in chans), reverse=True))

    def test_missing_source_comparison_yields_empty_channels_not_crash(self):
        self.v["chart_specs"] = [c for c in self.v["chart_specs"] if c["chart_id"] != "source_comparison"]
        dash = _executive_dashboard_data(self.folder, self.v)
        self.assertIsNotNone(dash)
        self.assertEqual(dash["channels"], [])

    def test_channel_score_matches_methodology_transformation(self):
        """Per-channel index == intelligence._brand_reputation on that channel's
        records — i.e. the same (weighted sentiment + 1)/2 × 100 as the headline."""
        def rec(platform, s, w):
            return {"platform": platform,
                    "ai_analysis": {"sentiment_score": s},
                    "intelligence": {"reputation_eligible": True, "reputation_weight": w,
                                     "impact_score": 0.5}}
        records = [rec("tiktok", 0.8, 1.0), rec("tiktok", -0.2, 1.0),
                   rec("news", 0.1, 2.0), rec("news", 0.5, 1.0)]
        breakdown = intel._source_breakdown(records)
        for platform in ("tiktok", "news"):
            rows = [r for r in records if r["platform"] == platform]
            self.assertEqual(breakdown[platform]["brand_reputation_index"],
                             intel._brand_reputation(rows)["index"])
        # and the dashboard shows exactly those values, rounded for display
        self.assertAlmostEqual(breakdown["tiktok"]["brand_reputation_index"], 65.0)
        self.assertAlmostEqual(breakdown["news"]["brand_reputation_index"], 61.67, places=2)


class DashboardChannelRenderTests(unittest.TestCase):
    def _dash(self, channels):
        return {"score": 58.5, "evidence_score": 70, "sentiment": [], "composition": [],
                "records_ready": 120, "voices": 60, "delta": 1.2,
                "followers_total": None, "views_total": None, "channels": channels}

    def test_strip_renders_every_channel_with_name_score_and_n(self):
        chans = [{"name": "x", "n": 38, "score": 58.0},
                 {"name": "tiktok", "n": 30, "score": 66.0},
                 {"name": "news", "n": 28, "score": 60.0}]
        _, slide = _blank_slide()
        _render_executive_dashboard(slide, self._dash(chans), "el", {"client": "OPAP", "topic": "EJ"})
        joined = " | ".join(_texts(slide))
        self.assertIn("BRAND REPUTATION ΑΝΑ ΚΑΝΑΛΙ", joined)
        for c in chans:
            self.assertIn(_upper_label(_channel_label(c["name"], "el")), joined)
            self.assertIn(f"{c['score']:.1f}  n={c['n']}", joined)

    def test_channel_without_score_shows_dash_but_keeps_its_n(self):
        chans = [{"name": "reddit", "n": 0, "score": None}]
        _, slide = _blank_slide()
        _render_executive_dashboard(slide, self._dash(chans), "en", {})
        joined = " | ".join(_texts(slide))
        self.assertIn("—  n=0", joined)

    def test_no_channels_no_strip(self):
        _, slide = _blank_slide()
        _render_executive_dashboard(slide, self._dash([]), "en", {})
        joined = " | ".join(_texts(slide))
        self.assertNotIn("BY CHANNEL", joined)

    def test_overflow_beyond_six_channels_collapses_into_more_cell(self):
        chans = [{"name": f"p{i}", "n": 30 - i, "score": 50.0 + i} for i in range(8)]
        _, slide = _blank_slide()
        _render_executive_dashboard(slide, self._dash(chans), "en", {})
        joined = " | ".join(_texts(slide))
        self.assertIn("+3", joined)
        self.assertIn("MORE CHANNELS", joined)

    def test_individual_comment_cards_never_carry_a_platform(self):
        entries = [{"author": "u1", "author_idx": 1, "platform": "tiktok", "date": "2/9",
                    "score": 0.9, "impact": 0.5, "text": "Σχόλιο χωρίς κανάλι."}]
        _, slide = _blank_slide()
        _render_top_comments(slide, entries, [], "positive", "el",
                             {"client": "OPAP", "topic": "EJ"}, {})
        joined = " | ".join(_texts(slide)).lower()
        self.assertNotIn("tiktok", joined)


class ContinuationSectionNumberTests(unittest.TestCase):
    def test_split_pages_share_one_section_number(self):
        entries = [{"author": f"u{i}", "author_idx": i + 1, "platform": "", "date": "2/9",
                    "score": round(0.9 - i * 0.05, 2), "impact": 0.5,
                    "text": "Κείμενο. " * 60} for i in range(10)]
        pplan = {
            "language": "el", "research_context": {"client": "OPAP", "topic": "EJ"},
            "top_comments": {"positive": entries, "negative": []},
            "slides": [
                {"slide_id": "cover", "slide_type": "cover", "title": "Cover", "notes": {}},
                {"slide_id": "s1", "slide_type": "generic", "title": "Πρώτο", "notes": {}},
                {"slide_id": "top_comments_positive", "slide_type": "top_comments_positive",
                 "title": "Top 10 Θετικά Σχόλια (1/2)",
                 "notes": {"part": 1, "parts": 2, "start": 0, "end": 5}},
                {"slide_id": "top_comments_positive_p2", "slide_type": "top_comments_positive",
                 "title": "Top 10 Θετικά Σχόλια (2/2)",
                 "notes": {"part": 2, "parts": 2, "start": 5, "end": 10}},
                {"slide_id": "s2", "slide_type": "generic", "title": "Μετά", "notes": {}},
            ],
        }
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "deck.pptx"
            generate_pptx(pplan, {"chart_specs": []}, out, Path(td) / "missing-logo.png")
            prs = Presentation(str(out))
            nums = []
            for slide in list(prs.slides)[1:]:
                nums.append(next(t for t in _texts(slide) if t in {"01", "02", "03", "04"}))
        # continuation shares 02; the slide after continues at 03 with no gap
        self.assertEqual(nums, ["01", "02", "02", "03"])


if __name__ == "__main__":
    unittest.main()
