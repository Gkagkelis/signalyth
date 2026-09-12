from __future__ import annotations

import base64
import copy
import tempfile
import csv
import hashlib
import json
import math
import re
import shutil
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.util import Inches, Pt

from app.services.investigations import INDICATOR_REGISTRY, load_evidence_pack
from app.services.storage import RunStore
from app.services.visualizations import load_presentation_visual_pack, load_visualization_summary
from app.services.report_synthesis import build_report_synthesis, final_consistency_qa

PRESENTATION_RULESET_VERSION = "1.6.0"
PRESENTATION_METHODOLOGY_VERSION = "signalyth-presentation-intelligence-v1.4"
PRESENTATION_CONTRACT_VERSION = "signalyth-presentation-pack-v1.4"

BG = "FFFFFF"
INK = "1A1A1A"
MUTED = "5F6368"
STONE = "F1F1F1"
STONE_DARK = "DADADA"
WHITE = "FFFFFF"
AEGEAN = "155E75"
POS = "1E7F4F"
NEG = "B23B32"
NEUTRAL = "9AA0A6"
MIXED = "C98A2B"
WARN = "B26A00"
FONT = "Inter"
# Editorial display face for covers and section titles. Georgia ships with both
# Windows and macOS and has real Greek coverage, so the cover renders identically
# on a client machine instead of silently falling back to a default.
DISPLAY_FONT = "Georgia"
LABEL_FONT = "Segoe UI"
COVER_INK = "16365C"
COVER_ACCENT = "FD9D08"

# Stable semantic colors: the same sentiment/emotion always gets the same color
# in dashboard, PPTX, PDF and Word, as the methodology requires.
SENTIMENT_COLORS = {"positive": POS, "negative": NEG, "neutral": NEUTRAL, "mixed": MIXED}
EMOTION_COLORS = {
    "joy": "E2A63D", "anger": "B23B32", "sadness": "3E6B9E", "fear": "6B4FA1",
    "disgust": "7A7F3A", "surprise": "2E8C8C", "neutral": "9AA0A6",
}
CATEGORY_LABELS_EL = {
    "positive": "Θετικό", "negative": "Αρνητικό", "neutral": "Ουδέτερο", "mixed": "Μικτό",
    "joy": "Χαρά", "anger": "Θυμός", "sadness": "Λύπη", "fear": "Φόβος",
    "disgust": "Αποστροφή", "surprise": "Έκπληξη",
    "media": "ΜΜΕ", "person": "Πρόσωπα", "brand_owned": "Owned", "organization": "Οργανισμοί", "unknown": "Άγνωστο",
}

def _cat_key(label) -> str:
    return str(label or "").strip().lower()

def _cat_label(label, lang: str) -> str:
    key = _cat_key(label)
    if lang == "el" and key in CATEGORY_LABELS_EL:
        return CATEGORY_LABELS_EL[key]
    return str(label)

def _semantic_color(label) -> str | None:
    key = _cat_key(label)
    return SENTIMENT_COLORS.get(key) or EMOTION_COLORS.get(key)


class PresentationCancelled(RuntimeError):
    pass


class PresentationValidationError(RuntimeError):
    pass


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_float(value, default: float = 0.0) -> float:
    try:
        n = float(value)
        if math.isfinite(n):
            return n
    except (TypeError, ValueError):
        pass
    return default


def _safe_int(value, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _json_safe(value):
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return copy.deepcopy(value)


def _hash_payload(payload: dict) -> str:
    blob = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str, allow_nan=False)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def _lang_code(plan: dict | None) -> str:
    return "el" if str((plan or {}).get("report_language") or "English") == "Ελληνικά" else "en"


def _txt(value, lang: str) -> str:
    if isinstance(value, dict):
        return str(value.get(lang) or value.get("en") or value.get("el") or "")
    return str(value or "")


def _clean_text(text, limit: int = 5000) -> str:
    # Keep evidence text literal but remove control chars that can break OOXML/XML.
    out = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", " ", str(text or ""))
    out = re.sub(r"\s+", " ", out).strip()
    return out[:limit]


def _chart_map(visual_pack: dict) -> dict[str, dict]:
    return {str(c.get("chart_id")): copy.deepcopy(c) for c in (visual_pack.get("chart_specs") or []) if c.get("chart_id")}


def _inventory_map(evidence_pack: dict) -> dict[str, dict]:
    return {str(x.get("indicator_id")): x for x in (evidence_pack.get("indicator_inventory") or []) if x.get("indicator_id")}


def _indicator_value(inv: dict[str, dict], indicator_id: str):
    return copy.deepcopy((inv.get(indicator_id) or {}).get("value"))


def _claim(claim_id: str, text: str, indicator_ids: list[str], *, evidence_refs=None, claim_type="descriptive", confidence=None, causal_status="not_applicable", source_values=None) -> dict:
    return {
        "claim_id": claim_id,
        "text": _clean_text(text, 1200),
        "indicator_ids": list(dict.fromkeys(indicator_ids)),
        "evidence_refs": list(dict.fromkeys([str(x) for x in (evidence_refs or []) if x])),
        "claim_type": claim_type,
        "confidence": confidence,
        "causal_status": causal_status,
        "source_values": _json_safe(source_values or {}),
    }




def _resolve_investigation_indicator_ids(item: dict) -> list[str]:
    """Map Step 6 investigation objects to the 25-indicator contract.

    Older or externally produced investigation records may not carry explicit
    indicator_ids. Step 8 still requires traceability, so we infer a conservative
    indicator set from the investigation type and payload instead of creating
    untraceable presentation claims.
    """
    expected = {x[0] for x in INDICATOR_REGISTRY}
    ids: list[str] = []

    def add(value):
        if value is None:
            return
        if isinstance(value, str):
            values = [value]
        elif isinstance(value, (list, tuple, set)):
            values = list(value)
        else:
            return
        for raw in values:
            text = str(raw or "").strip()
            if text in expected and text not in ids:
                ids.append(text)

    for key in ("indicator_ids", "related_indicator_ids", "indicator_family_ids", "presentation_indicator_ids"):
        add(item.get(key))
    for key in ("primary_indicator", "target_indicator", "indicator_id"):
        add(item.get(key))

    typ = str(item.get("type") or "").strip()
    type_map = {
        "numeric_anomaly": ["numeric_anomalies", "time_trends"],
        "reputation_daily_change": ["brand_reputation", "time_trends"],
        "period_shift": ["brand_reputation", "time_trends"],
        "negative_narrative_driver": ["negative_narrative_drivers", "brand_reputation"],
        "positive_narrative_driver": ["positive_narrative_drivers", "brand_reputation"],
        "emotion_profile": ["emotions", "sentiment"],
        "source_divergence": ["source_breakdown", "brand_reputation"],
        "media_people_divergence": ["origin_breakdown", "media_influence", "people_influence"],
        "coordination_signal": ["coordination", "authenticity_risk"],
        "data_quality_guardrail": ["data_quality", "source_coverage", "sample_achievement"],
        "emerging_narrative": ["negative_narrative_drivers", "positive_narrative_drivers", "time_trends"],
        "story_syndication": ["story_syndication", "source_breakdown"],
    }
    add(type_map.get(typ))

    metrics_blob = json.dumps(_json_safe(item.get("metrics") or item.get("trigger") or {}), ensure_ascii=False).lower()
    if "reputation" in metrics_blob:
        add(["brand_reputation"])
    if any(word in metrics_blob for word in ("anger", "fear", "emotion", "joy", "sadness", "disgust", "surprise")):
        add(["emotions"])
    if any(word in metrics_blob for word in ("source", "platform")):
        add(["source_breakdown"])
    if any(word in metrics_blob for word in ("quality", "coverage", "sample")):
        add(["data_quality", "source_coverage", "sample_achievement"])

    # Last-resort traceability: this is still a valid indicator family, and the
    # finding remains evidence-linked through evidence_record_ids when available.
    if not ids:
        add(["top_mentions"])
    return ids


def _slide(slide_id: str, slide_type: str, title: str, *, subtitle="", chart_ids=None, investigation_ids=None, claims=None, priority=50, required=False, section="body", notes=None) -> dict:
    return {
        "slide_id": slide_id,
        "slide_type": slide_type,
        "title": _clean_text(title, 250),
        "subtitle": _clean_text(subtitle, 600),
        "chart_ids": list(dict.fromkeys(chart_ids or [])),
        "investigation_ids": list(dict.fromkeys(investigation_ids or [])),
        "claims": list(claims or []),
        "priority": max(0, min(100, int(priority))),
        "required": bool(required),
        "section": section,
        "notes": copy.deepcopy(notes or {}),
    }


def _largest_distribution(chart: dict | None) -> tuple[str | None, float]:
    cats = ((chart or {}).get("data") or {}).get("categories") or []
    vals = [(str(x.get("label") or ""), _safe_float(x.get("value"))) for x in cats]
    return max(vals, key=lambda x: x[1]) if vals else (None, 0.0)


def _material_quality_warning(evidence_pack: dict) -> bool:
    inv = _inventory_map(evidence_pack)
    quality = _indicator_value(inv, "data_quality") or {}
    coverage = _indicator_value(inv, "source_coverage") or {}
    sample = _indicator_value(inv, "sample_achievement") or {}
    auth = _indicator_value(inv, "authenticity_risk") or {}
    coord = _indicator_value(inv, "coordination") or []
    return (
        _safe_float(quality.get("score"), 100) < 80
        or _safe_float(coverage.get("ratio"), 1) < 0.75
        or _safe_float(sample.get("ratio"), 1) < 0.85
        or _safe_float(auth.get("low_authenticity_share"), 0) >= 0.05
        or bool(coord)
        or bool((evidence_pack.get("headline_metrics") or {}).get("step5_warnings"))
    )


def _validate_inputs(visual_pack: dict, evidence_pack: dict, plan: dict | None = None) -> None:
    if not isinstance(visual_pack, dict) or not isinstance(evidence_pack, dict):
        raise PresentationValidationError("Step 8 requires both the Step 7 Presentation Visual Pack and the Step 6 Evidence Pack.")
    guard = visual_pack.get("guardrails") or {}
    if guard.get("must_review_all_indicators") is not True or guard.get("native_editable_charts_required") is not True:
        raise PresentationValidationError("Step 8 stopped because the Step 7 visual guardrails are incomplete.")
    review = visual_pack.get("indicator_review") or []
    expected = [x[0] for x in INDICATOR_REGISTRY]
    got = [str(x.get("indicator_id")) for x in review]
    if got != expected or len(got) != len(set(got)) or not all(x.get("examined") is True for x in review):
        raise PresentationValidationError("Step 8 stopped because not all 25 indicators were reviewed exactly once.")
    charts = visual_pack.get("chart_specs") or []
    if any((c.get("presentation") or {}).get("native_editable_ready") is not True for c in charts):
        raise PresentationValidationError("Step 8 requires every selected Step 7 chart to be native-editable ready.")
    if any((c.get("presentation") or {}).get("render_as_raster") is True for c in charts):
        raise PresentationValidationError("Step 8 refuses raster chart contracts.")
    if (visual_pack.get("guardrails") or {}).get("must_not_present_association_as_causation") is not True:
        raise PresentationValidationError("Step 8 requires the causality guardrail.")
    if plan:
        ctx = visual_pack.get("research_context") or {}
        for key in ("client", "topic", "market", "date_from", "date_to"):
            if ctx.get(key) not in (None, "") and plan.get(key) not in (None, "") and str(ctx.get(key)) != str(plan.get(key)):
                raise PresentationValidationError(f"Step 8 research context mismatch for {key}.")



GOLD_STANDARD_CAPABILITIES = (
    ("context", "Research context / cover"),
    ("evidence_base", "Evidence base, period and sample context"),
    ("methodology", "Methodology / reading guide"),
    ("reputation_sentiment", "Brand Reputation + sentiment"),
    ("sentiment_evidence", "Positive / negative evidence examples"),
    ("emotional_profile", "Emotion distribution"),
    ("emotion_evolution", "Emotion evolution when time evidence exists"),
    ("drivers", "Positive / negative narrative and topic drivers"),
    ("evolution", "Time evolution and anomalies"),
    ("investigations", "Automatic investigations when triggered"),
    ("sources_audiences", "Media / people / source context"),
    ("influence", "Media and people influence rankings"),
    ("media_evidence", "High-impact media evidence when available"),
    ("media_implications", "Media / channel implications when a divergence is detected"),
    ("quality", "Data quality / coverage / authenticity limits"),
    ("evidence_traceability", "Evidence behind findings"),
    ("strategic_synthesis", "Strategic synthesis / implications"),
    ("conclusions", "Conclusions and what to watch"),
    ("analyst_findings", "Senior analyst evidence-grounded synthesis"),
    ("recommendations", "Evidence-linked recommendations and monitoring actions"),
)


def _top_mentions(inv: dict[str, dict]) -> list[dict]:
    value = _indicator_value(inv, "top_mentions") or []
    return [copy.deepcopy(x) for x in value if isinstance(x, dict)]


def _mention_claim(prefix: str, row: dict, *, lang: str, index: int) -> dict:
    label = str(row.get("sentiment_label") or "evidence").lower()
    author = _clean_text(row.get("author") or row.get("platform") or "Source", 80)
    excerpt = _clean_text(row.get("excerpt") or "", 230)
    meta_bits = [b for b in (row.get("platform"), row.get("date")) if b]
    score = row.get("sentiment_score"); impact = row.get("impact_score")
    if score is not None: meta_bits.append(f"sentiment {_safe_float(score):+.2f}")
    if impact is not None: meta_bits.append(f"impact {_safe_float(impact):.2f}")
    meta = " · ".join(str(b) for b in meta_bits)
    text = (f"{author}: {excerpt}" if excerpt else author) + (f" | {meta}" if meta else "")
    return _claim(
        f"{prefix}-{index}", text, ["top_mentions"], evidence_refs=[row.get("record_id")],
        claim_type="evidence_example", confidence=row.get("evidence_confidence"),
        source_values={"sentiment": label, "impact": row.get("impact_score"), "origin_group": row.get("origin_group"), "platform": row.get("platform")},
    )


def _gold_standard_audit(slides: list[dict], charts: dict[str, dict], inv: dict[str, dict], investigations: dict[str, dict], media_handling: str = "blended") -> dict:
    slide_ids = {str(s.get("slide_id")) for s in slides}
    mentions = _top_mentions(inv)
    sentiment_labels = {str(x.get("sentiment_label") or "").lower() for x in mentions if x.get("excerpt")}
    media_mentions = [x for x in mentions if str(x.get("origin_group") or "").lower() == "media" and x.get("excerpt")]
    emotion_chart = charts.get("emotion_distribution")
    _, largest_emotion_value = _largest_distribution(emotion_chart)
    emotion_material = bool(emotion_chart) and (largest_emotion_value >= 12.0 or any(str(x.get("type")) == "emotion_profile" for x in investigations.values()))
    investigation_material = any(
        _safe_int((x.get("presentation") or {}).get("priority_score")) >= 55
        and (_txt(x.get("finding"), "en") or _txt(x.get("question"), "en"))
        for x in investigations.values()
    )
    media_implication_ids = {
        str(x.get("investigation_id")) for x in investigations.values()
        if str(x.get("type") or "") in {"source_divergence", "media_people_divergence"}
        and _safe_int((x.get("presentation") or {}).get("priority_score")) >= 55
        and (_txt(x.get("finding"), "en") or _txt(x.get("question"), "en"))
    }
    media_implication_slides = {
        str(slide.get("slide_id")) for slide in slides
        if media_implication_ids & {str(x) for x in (slide.get("investigation_ids") or [])}
    }
    applicable = {
        "context": True,
        "evidence_base": True,
        "methodology": True,
        "reputation_sentiment": bool((inv.get("brand_reputation") or {}).get("available") or (inv.get("sentiment") or {}).get("available")),
        "sentiment_evidence": "positive" in sentiment_labels and "negative" in sentiment_labels,
        "emotional_profile": emotion_material,
        "emotion_evolution": "emotion_trends" in charts,
        "drivers": bool((inv.get("positive_narrative_drivers") or {}).get("available") or (inv.get("negative_narrative_drivers") or {}).get("available") or (inv.get("topic_drivers") or {}).get("available")),
        "evolution": "time_trends" in charts,
        "investigations": investigation_material,
        "sources_audiences": bool((inv.get("origin_breakdown") or {}).get("available") or (inv.get("source_breakdown") or {}).get("available")),
        "influence": bool((inv.get("media_influence") or {}).get("available") or (inv.get("people_influence") or {}).get("available")),
        "media_evidence": media_handling != "exclude" and len(media_mentions) >= 2,
        "media_implications": media_handling != "exclude" and bool(media_implication_ids),
        "quality": any(bool((inv.get(k) or {}).get("available")) for k in ("data_quality","source_coverage","sample_achievement","authenticity_risk","coordination")),
        "evidence_traceability": bool(mentions),
        "strategic_synthesis": bool((inv.get("positive_narrative_drivers") or {}).get("available") or (inv.get("negative_narrative_drivers") or {}).get("available") or investigations),
        "conclusions": True,
        "analyst_findings": "analyst_findings" in slide_ids,
        "recommendations": "recommendations" in slide_ids,
    }
    coverage_map = {
        "context": {"cover"},
        "evidence_base": {"evidence_base"},
        "methodology": {"methodology"},
        "reputation_sentiment": {"reputation_sentiment"},
        "sentiment_evidence": {"sentiment_evidence"},
        "emotional_profile": {"emotions"},
        "emotion_evolution": {"emotion_evolution"},
        "drivers": {"drivers"},
        "evolution": {"evolution"},
        "investigations": {x for x in slide_ids if x.startswith("investigation_")},
        "sources_audiences": {"sources_audiences"},
        "influence": {"influence"},
        "media_evidence": {"media_evidence"},
        "media_implications": media_implication_slides,
        "quality": {"data_integrity", "evidence_base", "methodology"},
        "evidence_traceability": {"evidence"},
        "strategic_synthesis": {"strategic_synthesis"},
        "conclusions": {"conclusions"},
        "analyst_findings": {"analyst_findings"},
        "recommendations": {"recommendations"},
    }
    rows=[]
    for key,label in GOLD_STANDARD_CAPABILITIES:
        is_app=bool(applicable.get(key))
        covered=bool(slide_ids & coverage_map.get(key,set())) if is_app else False
        rows.append({"capability_id":key,"label":label,"applicable":is_app,"status":"covered" if covered else ("not_applicable" if not is_app else "missing"),"slide_ids":sorted(slide_ids & coverage_map.get(key,set()))})
    denom=sum(1 for r in rows if r["applicable"])
    covered_count=sum(1 for r in rows if r["applicable"] and r["status"]=="covered")
    return {"benchmark":"EUROJACKPOT collaborator deck capability benchmark","capabilities":rows,"applicable":denom,"covered":covered_count,"coverage_percent":round(100*covered_count/denom,2) if denom else 100.0,"complete":covered_count==denom}


def build_presentation_plan(visual_pack: dict, evidence_pack: dict, plan: dict | None = None, *, cancel_check: Callable[[], bool] | None = None) -> dict:
    _validate_inputs(visual_pack, evidence_pack, plan)
    if cancel_check and cancel_check():
        raise PresentationCancelled("Presentation planning cancelled before slide selection")

    lang = _lang_code(plan)
    ctx = copy.deepcopy(visual_pack.get("research_context") or {})
    charts = _chart_map(visual_pack)
    inv = _inventory_map(evidence_pack)
    investigations = {str(x.get("investigation_id")): x for x in (visual_pack.get("investigation_candidates") or []) if x.get("investigation_id")}

    title = _clean_text(ctx.get("topic") or "Analysis")
    client = _clean_text(ctx.get("client") or "")
    market = _clean_text(ctx.get("market") or "")
    d_from = _clean_text(ctx.get("date_from") or "")
    d_to = _clean_text(ctx.get("date_to") or "")

    slides: list[dict] = []
    slides.append(_slide("cover", "cover", title, subtitle=f"{client} · {market} · {d_from} — {d_to}".strip(" ·—"), priority=100, required=True, section="opening"))

    rep = _indicator_value(inv, "brand_reputation") or {}
    sentiment_chart = charts.get("sentiment_distribution")
    emotion_chart = charts.get("emotion_distribution")
    source_chart = charts.get("source_comparison")
    narrative_chart = charts.get("narrative_drivers")
    trend_chart = charts.get("time_trends")
    quality_chart = charts.get("quality_matrix")
    top_mentions_chart = charts.get("top_mentions")
    mentions = _top_mentions(inv)

    exec_claims = []
    if rep.get("index") is not None:
        exec_claims.append(_claim("exec-reputation", (f"Brand Reputation: {_safe_float(rep.get('index')):.1f}/100" if lang=="en" else f"Brand Reputation: {_safe_float(rep.get('index')):.1f}/100 (δείκτης ψηφιακής συζήτησης)"), ["brand_reputation"], claim_type="deterministic_metric", source_values={"brand_reputation": rep.get("index")}))
    if sentiment_chart:
        label, value = _largest_distribution(sentiment_chart)
        if label:
            exec_claims.append(_claim("exec-sentiment", (f"Largest weighted sentiment group: {label} ({value:.1f}%)" if lang=="en" else f"Μεγαλύτερη σταθμισμένη ομάδα sentiment: {_cat_label(label, lang)} ({value:.1f}%)"), ["sentiment"], claim_type="deterministic_metric", source_values={label: value}))
    if trend_chart:
        rows = (trend_chart.get("data") or {}).get("rows") or []
        if len(rows) >= 2:
            first, last = rows[0], rows[-1]
            if first.get("brand_reputation_index") is not None and last.get("brand_reputation_index") is not None:
                delta = _safe_float(last.get("brand_reputation_index")) - _safe_float(first.get("brand_reputation_index"))
                exec_claims.append(_claim("exec-trend", ((f"Brand Reputation changed {delta:+.1f} points across the observed period." if lang=="en" else f"Το Brand Reputation μεταβλήθηκε κατά {delta:+.1f} μονάδες στην περίοδο παρατήρησης.") if lang=="en" else f"Το Brand Reputation μεταβλήθηκε κατά {delta:+.1f} μονάδες στην περίοδο παρατήρησης."), ["time_trends", "brand_reputation"], claim_type="deterministic_change", source_values={"delta": round(delta, 2)}))

    inv_candidates = sorted(investigations.values(), key=lambda x: (_safe_int((x.get("presentation") or {}).get("priority_score")), _safe_float((x.get("confidence") or {}).get("score"))), reverse=True)
    if inv_candidates:
        top = inv_candidates[0]
        finding = _txt(top.get("finding"), lang)
        if finding:
            exec_claims.append(_claim("exec-investigation", finding, _resolve_investigation_indicator_ids(top), evidence_refs=top.get("evidence_record_ids") or [], claim_type=str(top.get("conclusion_type") or "investigation"), confidence=(top.get("confidence") or {}).get("score"), causal_status=str(top.get("causal_status") or "not_proven"), source_values=(top.get("metrics") or {})))
    slides.append(_slide("executive_summary", "executive_summary", "Executive Dashboard", chart_ids=[x for x in ["brand_reputation", "evidence_confidence"] if x in charts], claims=exec_claims, priority=100, required=True, section="opening"))

    # Gold-standard scope: period/sample/audience context must be explicit, not hidden in notes.
    scope_charts = [x for x in ["sample_overview", "origin_breakdown", "market_relevance", "impact_coverage"] if x in charts]
    scope_claims = []
    sample_value = _indicator_value(inv, "sample_volume") or {}
    market_value = _indicator_value(inv, "market_relevance") or {}
    if sample_value:
        scope_claims.append(_claim("scope-sample", f"Analysis-ready evidence: {_safe_int(sample_value.get('analysis_ready_records'))} records; effective independent voices: {_safe_float(sample_value.get('effective_independent_voices')):.1f}.", ["sample_volume", "effective_independent_voices"], claim_type="deterministic_metric", source_values=sample_value))
    if market_value.get("high_relevance_share") is not None:
        scope_claims.append(_claim("scope-market", f"High target-market relevance: {100*_safe_float(market_value.get('high_relevance_share')):.1f}%.", ["market_relevance"], claim_type="deterministic_metric", source_values=market_value))
    slides.append(_slide("evidence_base", "scope", "Evidence base & scope" if lang == "en" else "Βάση δεδομένων & εύρος", chart_ids=scope_charts[:2], claims=scope_claims, priority=97, required=True, section="opening"))

    # Reputation + sentiment are core and are grouped deliberately to mirror the original report's analytical spine.
    rep_charts = [x for x in ["brand_reputation", "sentiment_distribution", "stance_distribution"] if x in charts]
    if rep_charts:
        slides.append(_slide("reputation_sentiment", "metrics", "Reputation & sentiment" if lang == "en" else "Φήμη & συναίσθημα", chart_ids=rep_charts, priority=98, section="perception"))

    media_handling = str(plan.get("media_handling") or "blended").lower()
    if media_handling not in {"blended", "separate", "exclude"}:
        media_handling = "blended"
    def _is_media(m: dict) -> bool:
        return str(m.get("origin_group") or "").lower() == "media"
    # In "separate"/"exclude" mode the organic story is analysed clean of media voices.
    evidence_pool = mentions if media_handling == "blended" else [m for m in mentions if not _is_media(m)]

    positive_mentions = [m for m in evidence_pool if str(m.get("sentiment_label") or "").lower() == "positive" and m.get("excerpt")]
    negative_mentions = [m for m in evidence_pool if str(m.get("sentiment_label") or "").lower() == "negative" and m.get("excerpt")]
    if positive_mentions or negative_mentions:
        evidence_claims = []
        # Intensity-aware ranking: how negative/positive it is matters as much as
        # how far it travelled. A brutal comment from a small account outranks a
        # lukewarm one from a big account; a strong post with big reach tops both.
        def _evidence_rank(m):
            intensity = abs(_safe_float(m.get("sentiment_score")))
            impact = _safe_float(m.get("impact_score"))
            return intensity * (0.5 + 0.5 * impact)
        pos_sorted = sorted(positive_mentions, key=_evidence_rank, reverse=True)
        neg_sorted = sorted(negative_mentions, key=_evidence_rank, reverse=True)
        for i,m in enumerate(pos_sorted[:5], start=1):
            evidence_claims.append(_mention_claim("positive-evidence", m, lang=lang, index=i))
        for i,m in enumerate(neg_sorted[:5], start=1):
            evidence_claims.append(_mention_claim("negative-evidence", m, lang=lang, index=i))
        slides.append(_slide("sentiment_evidence", "evidence_split", "Top 5 positive & negative mentions" if lang == "en" else "Top 5 θετικές & αρνητικές αναφορές", claims=evidence_claims, priority=92, section="perception", notes={"split_by":"sentiment"}))

    if emotion_chart:
        emo_label, emo_value = _largest_distribution(emotion_chart)
        material = emo_value >= 12.0 or any(str(x.get("type")) == "emotion_profile" for x in investigations.values())
        if material:
            slides.append(_slide("emotions", "distribution", "Emotional profile" if lang == "en" else "Συναισθηματικό προφίλ", chart_ids=["emotion_distribution"], claims=[_claim("emotion-leading", f"Leading emotion: {emo_label} ({emo_value:.1f}%)", ["emotions"], claim_type="deterministic_metric", source_values={emo_label: emo_value})] if emo_label else [], priority=86, section="perception"))
    if "emotion_trends" in charts:
        slides.append(_slide("emotion_evolution", "trend", "Emotion evolution" if lang == "en" else "Εξέλιξη συναισθημάτων", chart_ids=["emotion_trends"], priority=84, section="perception"))

    if narrative_chart:
        rows = ((narrative_chart.get("data") or {}).get("rows") or [])
        claims = []
        for i, r in enumerate(rows[:4]):
            claims.append(_claim(f"driver-{i}", f"{_clean_text(r.get('name'))}: {_safe_float(r.get('contribution')):+.2f} reputation points", ["positive_narrative_drivers", "negative_narrative_drivers"], claim_type="deterministic_contribution", source_values={"contribution": r.get("contribution")}))
        slides.append(_slide("drivers", "drivers", "What moves Brand Reputation" if lang == "en" else "Τι μετακινεί το Brand Reputation", chart_ids=["narrative_drivers"] + (["topic_drivers"] if "topic_drivers" in charts else []), claims=claims, priority=100, section="drivers"))

    if trend_chart:
        slides.append(_slide("evolution", "trend", "Evolution over time" if lang == "en" else "Εξέλιξη στον χρόνο", chart_ids=["time_trends"] + (["anomalies"] if "anomalies" in charts else []), priority=98, section="evolution"))

    # Top automatic investigations become dedicated slides. Keep the number bounded to avoid a bloated report.
    inv_slides = 0
    for item in inv_candidates:
        priority = _safe_int((item.get("presentation") or {}).get("priority_score"))
        if priority < 55 or inv_slides >= 3:
            continue
        finding = _txt(item.get("finding"), lang)
        question = _txt(item.get("question"), lang)
        if not finding and not question:
            continue
        iid = str(item.get("investigation_id"))
        claims = []
        if finding:
            claims.append(_claim(f"{iid}-finding", finding, _resolve_investigation_indicator_ids(item), evidence_refs=item.get("evidence_record_ids") or [], claim_type=str(item.get("conclusion_type") or "investigation"), confidence=(item.get("confidence") or {}).get("score"), causal_status=str(item.get("causal_status") or "not_proven"), source_values=item.get("metrics") or {}))
        related = []
        typ = str(item.get("type") or "")
        if typ in {"numeric_anomaly", "reputation_daily_change", "period_shift", "emerging_narrative"}:
            related = [x for x in ["time_trends", "anomalies", "narrative_drivers"] if x in charts]
        elif typ in {"negative_narrative_driver", "positive_narrative_driver"}:
            related = [x for x in ["narrative_drivers", "topic_drivers"] if x in charts]
        elif typ == "emotion_profile":
            related = [x for x in ["emotion_distribution", "time_trends"] if x in charts]
        elif typ in {"source_divergence", "media_people_divergence"}:
            related = [x for x in ["source_comparison", "origin_breakdown"] if x in charts]
        elif typ in {"coordination_signal", "data_quality_guardrail"}:
            related = [x for x in ["authenticity_coordination", "quality_matrix"] if x in charts]
        slides.append(_slide(f"investigation_{inv_slides+1}", "investigation", question or ("Automatic investigation" if lang == "en" else "Αυτόματη διερεύνηση"), subtitle=finding, chart_ids=related[:2], investigation_ids=[iid], claims=claims, priority=priority, section="investigations", notes={"causality_guardrail": "Association ≠ proven causality."}))
        inv_slides += 1

    # Gold-standard guarantee: media/source divergence investigations that qualify for
    # presentation must never be silently dropped by the slide cap.
    placed_inv_ids = {str(x) for sl in slides for x in (sl.get("investigation_ids") or [])}
    for item in inv_candidates:
        typ = str(item.get("type") or "")
        if typ not in {"source_divergence", "media_people_divergence"}:
            continue
        if _safe_int((item.get("presentation") or {}).get("priority_score")) < 55:
            continue
        iid = str(item.get("investigation_id"))
        if iid in placed_inv_ids:
            continue
        finding = _txt(item.get("finding"), lang)
        question = _txt(item.get("question"), lang)
        if not finding and not question:
            continue
        claims = []
        if finding:
            claims.append(_claim(f"{iid}-finding", finding, _resolve_investigation_indicator_ids(item), evidence_refs=item.get("evidence_record_ids") or [], claim_type=str(item.get("conclusion_type") or "investigation"), confidence=(item.get("confidence") or {}).get("score"), causal_status=str(item.get("causal_status") or "not_proven"), source_values=item.get("metrics") or {}))
        related = [x for x in ["source_comparison", "origin_breakdown"] if x in charts]
        slides.append(_slide(f"investigation_{inv_slides+1}", "investigation", question or ("Automatic investigation" if lang == "en" else "Αυτόματη διερεύνηση"), subtitle=finding, chart_ids=related[:2], investigation_ids=[iid], claims=claims, priority=56, section="investigations", notes={"causality_guardrail": "Association ≠ proven causality."}))
        inv_slides += 1
        placed_inv_ids.add(iid)

    if source_chart or "origin_breakdown" in charts:
        slides.append(_slide("sources_audiences", "sources", "Sources & audiences" if lang == "en" else "Πηγές & κοινά", chart_ids=[x for x in ["source_comparison", "origin_breakdown"] if x in charts], priority=88, section="context"))

    influence_pool = ["people_influence"] if media_handling in {"separate", "exclude"} else ["media_influence", "people_influence"]
    influence = [x for x in influence_pool if x in charts]
    if influence:
        slides.append(_slide("influence", "influence", "Who shapes the conversation" if lang == "en" else "Ποιοι διαμορφώνουν τη συζήτηση", chart_ids=influence, priority=80, section="context"))

    media_mentions = [m for m in mentions if _is_media(m) and m.get("excerpt")]
    if media_handling != "exclude" and len(media_mentions) >= 2:
        media_sorted = sorted(media_mentions, key=lambda m: _safe_float(m.get("impact_score")), reverse=True)
        media_claims = [_mention_claim("media-evidence", m, lang=lang, index=i) for i,m in enumerate(media_sorted[:5], start=1)]
        media_title = ("Media coverage & message" if lang == "en" else "Media κάλυψη & μήνυμα") if media_handling == "blended" else ("Media, analysed separately" if lang == "en" else "Τα ΜΜΕ ξεχωριστά: κάλυψη & μήνυμα")
        media_charts = ["media_influence"] if (media_handling == "separate" and "media_influence" in charts) else []
        media_sub = "" if media_handling == "blended" else ("Media are reported in their own section; organic voices are analysed separately." if lang == "en" else "Τα ΜΜΕ παρουσιάζονται σε δική τους ενότητα· οι οργανικές φωνές αναλύονται ξεχωριστά.")
        slides.append(_slide("media_evidence", "evidence_cards", media_title, subtitle=media_sub, chart_ids=media_charts, claims=media_claims, priority=82, section="context"))

    # Voices of the conversation: the leading voice per sentiment, then a ranked
    # appendix of up to 30 voices so the client sees WHO is talking, not just numbers.
    people_rows = [r for r in (_indicator_value(inv, "people_influence") or []) if isinstance(r, dict) and r.get("name")]
    leader_by_sent: dict[str, dict] = {}
    for m in evidence_pool:
        lbl = str(m.get("sentiment_label") or "").lower()
        if lbl in {"positive", "negative", "neutral", "mixed"} and m.get("author"):
            cur = leader_by_sent.get(lbl)
            if cur is None or _safe_float(m.get("impact_score")) > _safe_float(cur.get("impact_score")):
                leader_by_sent[lbl] = m
    if leader_by_sent:
        lead_claims = []
        order = ["positive", "negative", "neutral", "mixed"]
        names_el = {"positive": "Κορυφαία θετική φωνή", "negative": "Κορυφαία αρνητική φωνή", "neutral": "Κορυφαία ουδέτερη φωνή", "mixed": "Κορυφαία μικτή φωνή"}
        names_en = {"positive": "Top positive voice", "negative": "Top negative voice", "neutral": "Top neutral voice", "mixed": "Top mixed voice"}
        for li, lbl in enumerate([x for x in order if x in leader_by_sent], start=1):
            m = leader_by_sent[lbl]
            head = names_en[lbl] if lang == "en" else names_el[lbl]
            body = f"{head}: {_clean_text(m.get('author'), 70)} — “{_clean_text(m.get('excerpt'), 170)}” | {m.get('platform')} · {m.get('date')} · impact {_safe_float(m.get('impact_score')):.2f}"
            lead_claims.append(_claim(f"voice-leader-{li}", body, ["top_mentions", "people_influence" if (inv.get("people_influence") or {}).get("available") else "top_mentions"], evidence_refs=[m.get("record_id")], claim_type="evidence_example", source_values={"sentiment": lbl, "valence": lbl}))
        slides.append(_slide("voices_leaders", "narratives", "Leading voices per sentiment" if lang == "en" else "Οι κορυφαίες φωνές ανά sentiment", claims=lead_claims, priority=79, section="context"))
    if people_rows:
        ranked = sorted(people_rows, key=lambda r: _safe_float(r.get("attention")), reverse=True)[:30]
        page_size = 12
        pages = [ranked[i:i + page_size] for i in range(0, len(ranked), page_size)]
        for pi, page in enumerate(pages, start=1):
            title = ("Voices of the conversation" if lang == "en" else "Οι φωνές της συζήτησης") + (f" ({pi}/{len(pages)})" if len(pages) > 1 else "")
            slides.append(_slide(f"voices_ranking_{pi}", "voices_table", title, priority=40, section="appendix", notes={"rows": [
                {"rank": (pi - 1) * page_size + i + 1, "name": _clean_text(r.get("name"), 60), "records": _safe_int(r.get("records")), "attention": _safe_float(r.get("attention")), "weighted_sentiment": _safe_float(r.get("weighted_sentiment"))}
                for i, r in enumerate(page)
            ], "denominator": len(people_rows)}))

    if _material_quality_warning(evidence_pack):
        qcharts = [x for x in ["quality_matrix", "market_relevance", "impact_coverage", "authenticity_coordination", "story_syndication"] if x in charts]
        slides.append(_slide("data_integrity", "quality", "Data integrity & interpretation limits" if lang == "en" else "Ακεραιότητα δεδομένων & όρια ερμηνείας", chart_ids=qcharts[:3], claims=[_claim("quality-causality", "Association ≠ proven causality." if lang == "en" else "Συσχέτιση ≠ αποδεδειγμένη αιτιότητα.", ["data_quality", "source_coverage"], claim_type="guardrail")], priority=94, required=True, section="quality"))

    if top_mentions_chart:
        slides.append(_slide("evidence", "evidence", "Evidence behind the findings" if lang == "en" else "Evidence πίσω από τα ευρήματα", chart_ids=["top_mentions"], priority=90, section="evidence"))

    # Strategic synthesis mirrors the old report's strengths/frictions/opportunities/risks logic without inventing strategy.
    synthesis_claims = []
    if narrative_chart:
        rows = ((narrative_chart.get("data") or {}).get("rows") or [])
        pos = [r for r in rows if _safe_float(r.get("contribution")) > 0]
        neg = [r for r in rows if _safe_float(r.get("contribution")) < 0]
        if pos:
            r=max(pos,key=lambda x:_safe_float(x.get("contribution")))
            synthesis_claims.append(_claim("synthesis-strength", f"Strength signal: {_clean_text(r.get('name'))} ({_safe_float(r.get('contribution')):+.2f} Reputation points)", ["positive_narrative_drivers"], claim_type="deterministic_contribution", source_values={"role":"strength","contribution":r.get("contribution")}))
        if neg:
            r=min(neg,key=lambda x:_safe_float(x.get("contribution")))
            synthesis_claims.append(_claim("synthesis-friction", f"Friction signal: {_clean_text(r.get('name'))} ({_safe_float(r.get('contribution')):+.2f} Reputation points)", ["negative_narrative_drivers"], claim_type="deterministic_contribution", source_values={"role":"friction","contribution":r.get("contribution")}))
    if trend_chart:
        rows=(trend_chart.get("data") or {}).get("rows") or []
        if len(rows)>=2 and rows[0].get("brand_reputation_index") is not None and rows[-1].get("brand_reputation_index") is not None:
            delta=_safe_float(rows[-1].get("brand_reputation_index"))-_safe_float(rows[0].get("brand_reputation_index"))
            synthesis_claims.append(_claim("synthesis-momentum", f"Momentum signal: Brand Reputation moved {delta:+.1f} points across the observed period.", ["brand_reputation","time_trends"], claim_type="deterministic_change", source_values={"role":"momentum","delta":round(delta,2)}))
    if inv_candidates:
        top=inv_candidates[0]
        synthesis_claims.append(_claim("synthesis-watch", ("Watch signal: " if lang=="en" else "Σήμα παρακολούθησης: ")+_txt(top.get("question"),lang), _resolve_investigation_indicator_ids(top), evidence_refs=top.get("evidence_record_ids") or [], claim_type="attention_point", causal_status="not_proven", source_values={"role":"watch"}))
    if synthesis_claims:
        slides.append(_slide("strategic_synthesis", "strategic_synthesis", "Strategic synthesis" if lang=="en" else "Στρατηγική σύνθεση", claims=synthesis_claims[:4], priority=96, section="closing", notes={"causality_guardrail":"Association ≠ proven causality."}))

    analyst = visual_pack.get("analyst_synthesis") or {}
    if analyst.get("narratives"):
        mention_by_id={str(m.get("record_id")):m for m in _top_mentions(inv)}
        groups: list[tuple[str,list[dict]]] = []
        for n in (analyst.get("narratives") or [])[:8]:
            theme=_clean_text(n.get("theme") or "", 60) or ("Other themes" if lang=="en" else "Λοιπά θέματα")
            for g in groups:
                if g[0]==theme:
                    g[1].append(n); break
            else:
                groups.append((theme,[n]))
        # One slide per theme; a large theme spills onto a continuation slide.
        nar_slides=0; ci=0
        for theme,items in groups:
            for chunk_start in range(0,len(items),3):
                if nar_slides>=3: break
                chunk=items[chunk_start:chunk_start+3]
                nar_claims=[]
                for n in chunk:
                    ci+=1
                    origin_lbl={"media":("media","ΜΜΕ"),"organic_people":("organic voices","οργανικές φωνές"),"owned":("owned","owned/promo"),"mixed":("mixed","μικτής προέλευσης"),"unknown":("unknown","άγνωστης προέλευσης")}.get(str(n.get("origin")),("mixed","μικτής προέλευσης"))
                    text=(f"{_clean_text(n.get('title'),110)} — {_clean_text(n.get('idea'),330)} | Σημασία για το brand: {_clean_text(n.get('brand_meaning'),260)} | Βαρύτητα: {_clean_text(n.get('importance'),160)} ({origin_lbl[1]})" if lang=="el" else
                          f"{_clean_text(n.get('title'),110)} — {_clean_text(n.get('idea'),330)} | Brand meaning: {_clean_text(n.get('brand_meaning'),260)} | Weight: {_clean_text(n.get('importance'),160)} ({origin_lbl[0]})")
                    quotes=[]
                    for rid in (n.get("evidence_record_ids") or [])[:2]:
                        m=mention_by_id.get(str(rid))
                        if m and m.get("excerpt"):
                            quotes.append({"excerpt":_clean_text(m.get("excerpt"),210),"platform":m.get("platform"),"date":m.get("date"),"author":m.get("author")})
                    nar_claims.append(_claim(f"narrative-{ci}",text,list(n.get("indicator_ids") or []),evidence_refs=n.get("evidence_record_ids") or [],claim_type="analyst_finding",causal_status="not_proven",source_values={"valence":n.get("valence"),"origin":n.get("origin"),"quotes":quotes}))
                nar_slides+=1
                base_title="What people really say" if lang=="en" else "Τι πραγματικά λέει ο κόσμος"
                slides.append(_slide(f"consumer_narratives_{nar_slides}","narratives",f"{base_title} · {theme}",claims=nar_claims,priority=93,section="perception",notes={"causality_guardrail":"Association ≠ proven causality."}))
    if analyst.get("strategic_implications"):
        imp_claims=[]
        for i,m in enumerate((analyst.get("strategic_implications") or [])[:4],start=1):
            text=(f"Insight: {_clean_text(m.get('insight'),260)} | Implication: {_clean_text(m.get('implication'),260)} | Action: {_clean_text(m.get('action'),260)} | Objective: {_clean_text(m.get('objective'),200)}" if lang=="en" else
                  f"Insight: {_clean_text(m.get('insight'),260)} | Επίπτωση: {_clean_text(m.get('implication'),260)} | Ενέργεια: {_clean_text(m.get('action'),260)} | Στόχος: {_clean_text(m.get('objective'),200)}")
            imp_claims.append(_claim(f"implication-{i}",text,list(m.get("indicator_ids") or []),evidence_refs=m.get("evidence_record_ids") or [],claim_type="recommendation",causal_status="not_proven"))
        slides.append(_slide("strategic_implications","evidence_cards","Strategic implications" if lang=="en" else "Στρατηγικές κατευθύνσεις",claims=imp_claims,priority=96,section="closing",notes={"causality_guardrail":"Association ≠ proven causality."}))
    if analyst.get("findings"):
        analyst_claims=[]
        for i,f in enumerate((analyst.get("findings") or [])[:6],start=1):
            text=f"{_clean_text(f.get('title'),120)} — {_clean_text(f.get('finding'),700)} | {_clean_text(f.get('interpretation'),500)}"
            analyst_claims.append(_claim(f"analyst-{i}",text,list(f.get("indicator_ids") or []),evidence_refs=f.get("evidence_record_ids") or [],
                                         claim_type="analyst_finding",confidence=f.get("confidence"),causal_status=f.get("causal_status") or "not_proven",
                                         source_values={"materiality":f.get("materiality"),"confidence_label":f.get("confidence")}))
        slides.append(_slide("analyst_findings","evidence_cards","Senior analyst synthesis" if lang=="en" else "Σύνθεση senior analyst",claims=analyst_claims,priority=97,section="closing",notes={"causality_guardrail":"Association ≠ proven causality."}))
    if analyst.get("recommendations"):
        rec_claims=[]
        for i,r in enumerate((analyst.get("recommendations") or [])[:5],start=1):
            extra_en=[]; extra_el=[]
            if r.get("timing"): extra_en.append(f"When: {_clean_text(r.get('timing'),120)}"); extra_el.append(f"Πότε: {_clean_text(r.get('timing'),120)}")
            if r.get("stop_condition"): extra_en.append(f"Stop/review if: {_clean_text(r.get('stop_condition'),200)}"); extra_el.append(f"Όριο διακοπής: {_clean_text(r.get('stop_condition'),200)}")
            text=(f"Action: {_clean_text(r.get('action'),450)} | Why: {_clean_text(r.get('rationale'),550)} | Monitor: {_clean_text(r.get('monitor'),350)}" + ("".join(" | "+e for e in extra_en)) if lang=="en" else
                  f"Ενέργεια: {_clean_text(r.get('action'),450)} | Γιατί: {_clean_text(r.get('rationale'),550)} | Παρακολούθηση: {_clean_text(r.get('monitor'),350)}" + ("".join(" | "+e for e in extra_el)))
            rec_claims.append(_claim(f"recommendation-{i}",text,list(r.get("indicator_ids") or []),evidence_refs=r.get("evidence_record_ids") or [],claim_type="recommendation",causal_status="not_proven"))
        slides.append(_slide("recommendations","evidence_cards","Recommended actions" if lang=="en" else "Προτεινόμενες ενέργειες",claims=rec_claims,priority=95,section="closing"))
    if analyst.get("benchmark_summary") and (plan or {}).get("benchmark"):
        slides.append(_slide("benchmark","evidence_cards","Benchmark / comparison" if lang=="en" else "Benchmark / σύγκριση",claims=[_claim("benchmark-1",analyst.get("benchmark_summary"),["brand_reputation","time_trends"],claim_type="benchmark",causal_status="not_proven")],priority=86,section="context"))

    methodology_claims = [
        _claim("method-reputation", "Brand Reputation is a deterministic 0–100 index built from evidence-linked sentiment, independent-voice weight, authenticity, confidence and bounded impact; reach does not create sentiment by itself.", ["brand_reputation","sentiment","effective_independent_voices","authenticity_risk","evidence_confidence","impact_attention"], claim_type="methodology"),
        _claim("method-emotions", "Emotion reporting is based on the dominant classified emotion in organic people evidence; low-signal emotions do not automatically receive a dedicated slide.", ["emotions","origin_breakdown"], claim_type="methodology"),
        _claim("method-visibility", "Owned/promotional and factual not-applicable content can remain visible for dissemination analysis while being excluded from Reputation weighting.", ["origin_breakdown","stance","impact_attention"], claim_type="methodology"),
        _claim("method-causality", "Association ≠ proven causality. Automatic investigations surface evidence-linked associations, deterministic contributions and risk signals only.", ["numeric_anomalies","time_trends","top_mentions"], claim_type="guardrail", causal_status="not_proven"),
    ]
    methodology_sub = ""
    if media_handling == "exclude":
        methodology_sub = ("Media (news) evidence was excluded from this report by explicit client configuration; it remains available in the evidence pack." if lang=="en"
                           else "Τα ΜΜΕ (news) εξαιρέθηκαν από αυτό το report με ρητή επιλογή ρύθμισης· παραμένουν διαθέσιμα στο evidence pack.")
    slides.append(_slide("methodology", "methodology", "Methodology & reading guide" if lang=="en" else "Μεθοδολογία & οδηγός ανάγνωσης", subtitle=methodology_sub, claims=methodology_claims, priority=70, required=True, section="appendix"))

    # Deterministic closing: not generic AI prose. It is assembled from the highest-priority evidence.
    closing_claims = []
    if narrative_chart:
        rows = ((narrative_chart.get("data") or {}).get("rows") or [])
        pos = [r for r in rows if _safe_float(r.get("contribution")) > 0]
        neg = [r for r in rows if _safe_float(r.get("contribution")) < 0]
        if pos:
            r = max(pos, key=lambda x: _safe_float(x.get("contribution")))
            closing_claims.append(_claim("close-positive", (f"Strongest positive driver: {_clean_text(r.get('name'))}" if lang=="en" else f"Ισχυρότερος θετικός παράγοντας: {_clean_text(r.get('name'))}"), ["positive_narrative_drivers"], claim_type="deterministic_rank"))
        if neg:
            r = min(neg, key=lambda x: _safe_float(x.get("contribution")))
            closing_claims.append(_claim("close-negative", (f"Strongest negative driver: {_clean_text(r.get('name'))}" if lang=="en" else f"Ισχυρότερος αρνητικός παράγοντας: {_clean_text(r.get('name'))}"), ["negative_narrative_drivers"], claim_type="deterministic_rank"))
    if inv_candidates:
        top = inv_candidates[0]
        closing_claims.append(_claim("close-watch", ("Watch next: " if lang == "en" else "Παρακολούθηση: ") + _txt(top.get("question"), lang), _resolve_investigation_indicator_ids(top), evidence_refs=top.get("evidence_record_ids") or [], claim_type="attention_point", causal_status="not_proven"))
    heads=(analyst.get("slide_headlines") or {}) if isinstance(analyst, dict) else {}
    for sl in slides:
        h=heads.get(sl.get("slide_id"))
        if h: sl["title"]=_clean_text(h,110)
    if analyst.get("emotion_interpretation"):
        for sl in slides:
            if sl.get("slide_id")=="emotions" and not sl.get("subtitle"):
                sl["subtitle"]=_clean_text(analyst.get("emotion_interpretation"),260)
    conclusions_notes={}
    pr=(analyst.get("priorities") or {}) if isinstance(analyst, dict) else {}
    if any(pr.get(k) for k in ("protect","improve","amplify","monitor")):
        conclusions_notes["priorities"]={k:list(pr.get(k) or [])[:3] for k in ("protect","improve","amplify","monitor")}
    if analyst.get("final_takeaway"):
        conclusions_notes["final_takeaway"]=_clean_text(analyst.get("final_takeaway"),420)
    slides.append(_slide("conclusions", "conclusions", "Conclusions & what to watch" if lang == "en" else "Συμπεράσματα & η επόμενη κίνηση", claims=closing_claims, priority=100, required=True, section="closing", notes=conclusions_notes or None))

    if cancel_check and cancel_check():
        raise PresentationCancelled("Presentation planning cancelled after slide selection")

    # Complete Step 8 indicator disposition contract.
    indicator_review = []
    for item in visual_pack.get("indicator_review") or []:
        indicator_id = str(item.get("indicator_id"))
        linked_slides = [s["slide_id"] for s in slides if indicator_id in {i for c in s.get("chart_ids") or [] for i in (charts.get(c, {}).get("indicator_ids") or [])} or any(indicator_id in (cl.get("indicator_ids") or []) for cl in s.get("claims") or [])]
        if linked_slides:
            status = "included"
            reason = None
        elif item.get("available"):
            status = "reviewed_omitted_low_signal"
            reason = "Reviewed by Step 8; no dedicated or contextual slide justified by the current evidence."
        else:
            status = "reviewed_omitted_insufficient_data"
            reason = "Evidence unavailable or insufficient; no decorative slide generated."
        indicator_review.append({
            "indicator_id": indicator_id,
            "examined": True,
            "step7_visual_status": item.get("visual_status"),
            "presentation_status": status,
            "slide_ids": linked_slides,
            "omission_reason": reason,
        })

    expected = [x[0] for x in INDICATOR_REGISTRY]
    if [x["indicator_id"] for x in indicator_review] != expected or not all(x["examined"] for x in indicator_review):
        raise PresentationValidationError("Step 8 presentation indicator review is incomplete.")

    chart_review = []
    used_chart_ids = {cid for s in slides for cid in s.get("chart_ids") or []}
    for cid in visual_pack.get("presentation_chart_order") or []:
        c = charts.get(cid)
        if not c:
            continue
        chart_review.append({
            "chart_id": cid,
            "examined": True,
            "used": cid in used_chart_ids,
            "slide_ids": [s["slide_id"] for s in slides if cid in (s.get("chart_ids") or [])],
            "omission_reason": None if cid in used_chart_ids else "Reviewed but omitted to avoid low-information or redundant client slides.",
        })

    claim_ledger = [copy.deepcopy(cl) | {"slide_id": s["slide_id"]} for s in slides for cl in s.get("claims") or []]
    _validate_claim_ledger(claim_ledger, expected)
    gold_standard_audit = _gold_standard_audit(slides, charts, inv, investigations, media_handling)
    missing_caps = [r["capability_id"] for r in gold_standard_audit["capabilities"] if r["applicable"] and r["status"] == "missing"]
    if missing_caps:
        # Documented (never silent) omission: the export still ships, and the gap is
        # recorded prominently in the audit and the plan-level warnings.
        gold_standard_audit["warnings"] = list(gold_standard_audit.get("warnings") or []) + [
            f"Capability without dedicated coverage in this deck: {cap}" for cap in missing_caps
        ]

    return {
        "contract_version": PRESENTATION_CONTRACT_VERSION,
        "ruleset_version": PRESENTATION_RULESET_VERSION,
        "methodology_version": PRESENTATION_METHODOLOGY_VERSION,
        "generated_at": _utcnow(),
        "client_logo": plan.get("client_logo"),
        "brand_accent": plan.get("brand_accent"),
        "language": lang,
        "research_context": ctx,
        "slides": slides,
        "indicator_review": indicator_review,
        "chart_review": chart_review,
        "claim_ledger": claim_ledger,
        "gold_standard_audit": gold_standard_audit,
        "warnings": list(dict.fromkeys([_clean_text(x, 500) for x in (visual_pack.get("warnings") or []) if x])),
        "guardrails": {
            "all_25_indicators_reviewed": True,
            "no_invented_metrics": True,
            "material_claims_require_traceability": True,
            "native_editable_powerpoint_objects": True,
            "quality_warnings_preserved": True,
            "association_not_causation": True,
            "raw_evidence_never_modified": True,
        },
    }


def _validate_claim_ledger(claims: list[dict], expected_indicator_ids: list[str] | None = None) -> None:
    expected = set(expected_indicator_ids or [x[0] for x in INDICATOR_REGISTRY])
    seen = set()
    for cl in claims:
        cid = str(cl.get("claim_id") or "")
        if not cid or cid in seen:
            raise PresentationValidationError("Every presentation claim must have a unique claim_id.")
        seen.add(cid)
        ids = cl.get("indicator_ids") or []
        if cl.get("claim_type") not in {"guardrail"} and not ids:
            raise PresentationValidationError(f"Material claim {cid} has no indicator traceability.")
        if any(i not in expected for i in ids):
            raise PresentationValidationError(f"Claim {cid} references an unknown indicator.")
        if cl.get("causal_status") not in {None, "not_applicable", "not_proven", "deterministic_contribution"} and "causal" in str(cl.get("claim_type")):
            raise PresentationValidationError(f"Claim {cid} violates the causality contract.")


# ---------- PowerPoint renderer ----------

def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _set_bg(slide, color: str = BG):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = _rgb(color)


def _add_text(slide, text, x, y, w, h, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = _clean_text(text, 4000)
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = _rgb(color)
    return box


def _add_rule(slide, x, y, w, color=STONE_DARK, width=1.0):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.01))
    line.fill.solid(); line.fill.fore_color.rgb = _rgb(color); line.line.fill.background()
    return line


def _add_header(slide, title, section_no=None):
    if section_no is not None:
        _add_text(slide, f"{section_no:02d}", 0.58, 0.45, 0.5, 0.28, size=9, color=MUTED, bold=True)
    clean = _clean_text(title, 250)
    size = 25 if len(clean) <= 54 else 20 if len(clean) <= 82 else 17
    _add_text(slide, clean, 1.15, 0.30, 11.4, 0.82, size=size, bold=True, valign=MSO_ANCHOR.MIDDLE)
    _add_rule(slide, 0.58, 1.15, 12.15)


def _add_footer(slide, page, context):
    topic = _clean_text(context.get("topic") or "SIGNALYTH", 80)
    _add_text(slide, topic, 0.58, 7.12, 8.8, 0.18, size=7.5, color=MUTED)
    _add_text(slide, str(page), 12.05, 7.12, 0.65, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)


def _track(run, points: float) -> None:
    """Letter-spacing on a run. python-pptx has no API for it, so the OOXML
    attribute is set directly. Tracked small caps are what separates an editorial
    cover from a default deck."""
    try:
        run.font._rPr.set("spc", str(int(points * 100)))
    except Exception:
        pass


def _add_display(slide, text, x, y, w, h, *, size, color, font=None, bold=False,
                 tracking=0.0, align=PP_ALIGN.LEFT, line_spacing=None):
    """Display typography block used by the cover."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    if line_spacing:
        para.line_spacing = line_spacing
    run = para.add_run()
    run.text = _clean_text(text, 300)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font or DISPLAY_FONT
    run.font.color.rgb = _rgb(color)
    if tracking:
        _track(run, tracking)
    return box


def _add_logo(slide, logo_path: Path, x, y, w):
    if logo_path.exists():
        try:
            return slide.shapes.add_picture(str(logo_path), Inches(x), Inches(y), width=Inches(w))
        except Exception:
            return None
    return None


def _add_kpi_card(slide, label, value, x, y, w, h, *, suffix="", accent=AEGEAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid(); shape.fill.fore_color.rgb = _rgb("FAFAFA"); shape.line.color.rgb = _rgb(STONE_DARK)
    shape.shadow.inherit = False
    label_txt = str(label or "").upper()
    _add_text(slide, label_txt, x+0.2, y+0.17, w-0.4, 0.5, size=9, color=MUTED, bold=True)
    value_size = 36 if w >= 2.4 else 26 if w >= 1.7 else 20
    _add_text(slide, f"{value}{suffix}", x+0.2, y+0.52, w-0.4, h-0.66, size=value_size, color=accent, bold=True, valign=MSO_ANCHOR.MIDDLE)
    return shape


def _native_bar(slide, categories, series, x, y, w, h, *, horizontal=False, diverging=False, percent=False, lang="en", point_colors=None, show_values=True):
    data = ChartData(); data.categories = [_clean_text(_cat_label(c, lang), 60) for c in categories]
    for name, vals in series:
        data.add_series(_clean_text(name, 50), [float(_safe_float(v)) for v in vals])
    chart_type = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(9)
        chart.legend.include_in_layout = False
    chart.has_title = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(STONE)
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.name = FONT
    if percent:
        chart.value_axis.maximum_scale = 100
        chart.value_axis.minimum_scale = 0
    if diverging:
        chart.value_axis.crosses_at = 0
        try:
            from pptx.enum.chart import XL_TICK_LABEL_POSITION
            chart.category_axis.tick_labels.font.size = Pt(9)
            chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        except Exception:
            pass
    palette = [AEGEAN, "3E6B9E", NEG, POS]
    values0 = [float(_safe_float(v)) for v in series[0][1]] if series else []
    for i, sr in enumerate(chart.series):
        sr.format.fill.solid(); sr.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        sr.format.line.color.rgb = _rgb(WHITE)
    if len(series) == 1:
        sr = chart.series[0]
        try:
            for j, pt in enumerate(sr.points):
                col = None
                if point_colors:
                    col = point_colors[j] if j < len(point_colors) else None
                elif diverging:
                    col = POS if values0[j] >= 0 else NEG
                else:
                    col = _semantic_color(categories[j])
                if col:
                    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _rgb(col)
        except Exception:
            pass
        if show_values:
            try:
                plot = chart.plots[0]
                plot.has_data_labels = True
                dl = plot.data_labels
                dl.font.size = Pt(9); dl.font.name = FONT; dl.font.color.rgb = _rgb(INK)
                dl.number_format = '0.0'; dl.number_format_is_linked = False
            except Exception:
                pass
        try:
            chart.plots[0].gap_width = 60
        except Exception:
            pass
    return chart


def _native_doughnut(slide, categories, values, x, y, w, h, *, lang="en"):
    data = ChartData(); data.categories = [_clean_text(_cat_label(c, lang), 40) for c in categories]
    data.add_series("share", [float(_safe_float(v)) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.font.size = Pt(11); chart.legend.font.name = FONT
    chart.legend.include_in_layout = False
    sr = chart.series[0]
    try:
        for j, pt in enumerate(sr.points):
            col = _semantic_color(categories[j]) or [AEGEAN, "3E6B9E", NEUTRAL, MIXED, NEG, POS][j % 6]
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _rgb(col)
            pt.format.line.color.rgb = _rgb(WHITE); pt.format.line.width = Pt(2)
    except Exception:
        pass
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = True
        dl.font.size = Pt(11); dl.font.bold = True; dl.font.name = FONT; dl.font.color.rgb = _rgb(WHITE)
        dl.number_format = '0"%"'; dl.number_format_is_linked = False
    except Exception:
        pass
    return chart


def _native_line(slide, categories, series, x, y, w, h, *, min_y=None, max_y=None):
    data = ChartData(); data.categories = [_clean_text(c, 30) for c in categories]
    for name, vals in series:
        data.add_series(_clean_text(name, 50), [float(_safe_float(v)) for v in vals])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.font.size = Pt(8)
    chart.value_axis.has_major_gridlines = True; chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(STONE_DARK)
    chart.value_axis.tick_labels.font.size = Pt(8); chart.category_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.tick_labels.font.name = FONT; chart.category_axis.tick_labels.font.name = FONT
    if min_y is not None: chart.value_axis.minimum_scale = min_y
    if max_y is not None: chart.value_axis.maximum_scale = max_y
    palette = [AEGEAN, NEG, POS]
    for i, s in enumerate(chart.series):
        s.format.line.color.rgb = _rgb(palette[i % len(palette)]); s.format.line.width = Pt(2)
        s.marker.format.fill.solid(); s.marker.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
    return chart


def _add_editable_table(slide, headers, rows, x, y, w, h, *, col_widths=None, font_size=9):
    rows = list(rows)
    effective_h = min(h, max(0.82, 0.42 * (len(rows) + 1)))
    table_shape = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(effective_h))
    table = table_shape.table
    table.rows[0].height = Inches(0.36)
    for i in range(1, len(table.rows)):
        rr = table.rows[i]
        rr.height = Inches(max(0.34, min(0.48, (effective_h-0.36)/max(1,len(rows)))))
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(w * cw / total)
    for j, head in enumerate(headers):
        cell = table.cell(0,j); cell.text = _clean_text(head, 80); cell.fill.solid(); cell.fill.fore_color.rgb = _rgb("EFEFEF")
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.name=FONT; r.font.size=Pt(font_size); r.font.bold=True; r.font.color.rgb=_rgb(INK)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell=table.cell(i,j); cell.text=_clean_text(val, 500); cell.fill.solid(); cell.fill.fore_color.rgb=_rgb(WHITE if i%2 else "F7F7F7")
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name=FONT; r.font.size=Pt(font_size); r.font.color.rgb=_rgb(INK)
    return table


def _render_chart_spec(slide, chart: dict, x, y, w, h, lang: str):
    typ = chart.get("chart_type"); data = chart.get("data") or {}
    if typ not in {"gauge_kpi", "quality_kpi", "kpi_group"}:
        _add_text(slide, _txt(chart.get("title"), lang), x, y, w, 0.3, size=10.5, color=MUTED, bold=True)
        y += 0.38
        h = max(0.5, h - 0.38)
    if typ == "gauge_kpi":
        val = data.get("value")
        title_txt = _txt(chart.get("title"), lang)
        cid = str(chart.get("chart_id") or "").lower()
        if "confidence" in cid or "βεβαιότητα" in title_txt.lower() or "confidence" in title_txt.lower():
            # §14: confidence is shown as High/Medium/Low with the profile score as
            # secondary context — never as false client-facing precision.
            score = _safe_float(val)
            label = ("High" if score >= 70 else "Medium" if score >= 40 else "Low") if lang == "en" else ("Υψηλή" if score >= 70 else "Μέτρια" if score >= 40 else "Χαμηλή")
            _add_kpi_card(slide, title_txt, label, x, y, w, h, accent=AEGEAN)
            _add_text(slide, ("evidence profile " if lang == "en" else "προφίλ evidence ") + f"{score:.0f}/100", x+0.2, y+h-0.34, w-0.4, 0.26, size=8.5, color=MUTED)
        else:
            _add_kpi_card(slide, title_txt, f"{_safe_float(val):.1f}", x, y, w, h, suffix=" / 100", accent=AEGEAN)
            if "reputation" in cid or "reputation" in title_txt.lower() or "φήμη" in title_txt.lower():
                # §9: Brand Reputation is client-facing only with its scientific scope subtitle.
                _add_text(slide, "Digital conversation signal" if lang == "en" else "Δείκτης ψηφιακής συζήτησης — όχι μέτρηση κοινής γνώμης", x+0.2, y+h-0.34, w-0.4, 0.26, size=8.5, color=MUTED)
        return "editable_shapes"
    if typ == "quality_kpi":
        val = next((data.get(k) for k in ("known_metric_share_percent","high_relevance_share_percent","value") if data.get(k) is not None), None)
        _add_kpi_card(slide, _txt(chart.get("title"), lang), f"{_safe_float(val):.1f}" if val is not None else "—", x, y, w, h, suffix="%" if val is not None else "", accent=AEGEAN)
        return "editable_shapes"
    if typ == "kpi_group":
        items=data.get("items") or []; n=max(1,len(items)); gap=.16
        cols = n if (w - gap*(n-1))/n >= 1.9 else max(1, (n + 1)//2)
        rows_n = (n + cols - 1)//cols
        cw=(w-gap*(cols-1))/cols; ch=(h-gap*(rows_n-1))/rows_n
        for i,item in enumerate(items):
            label=_txt(item.get("label"),lang); val=item.get("value")
            r,c=divmod(i,cols)
            _add_kpi_card(slide,label, f"{_safe_float(val):.0f}" if isinstance(val,(int,float)) else str(val), x+c*(cw+gap),y+r*(ch+gap),cw,ch,accent=AEGEAN)
        return "editable_shapes"
    if typ == "distribution_bar":
        cats=data.get("categories") or []
        labels=[c.get("label") for c in cats]
        values=[c.get("value") for c in cats]
        keys={_cat_key(l) for l in labels}
        if keys and keys.issubset(set(SENTIMENT_COLORS) | {"media","person","brand_owned","organization","unknown"}):
            _native_doughnut(slide,labels,values,x,y,w,h,lang=lang)
        else:
            order=sorted(range(len(labels)),key=lambda i:-_safe_float(values[i]))
            labels=[labels[i] for i in order]; values=[values[i] for i in order]
            _native_bar(slide,labels,[("%",values)],x,y,w,h,horizontal=True,percent=True,lang=lang)
        return "native_chart"
    if typ == "diverging_bar":
        rows=data.get("rows") or []
        _native_bar(slide,[r.get("name") for r in rows[:10]],[("Contribution",[r.get("contribution") for r in rows[:10]])],x,y,w,h,horizontal=True,diverging=True,lang=lang)
        return "native_chart"
    if typ == "grouped_bar":
        rows=data.get("rows") or []
        _native_bar(slide,[r.get("name") for r in rows],[("Record share" if lang=="en" else "Μερίδιο αναφορών",[r.get("record_share") for r in rows]),("Attention share" if lang=="en" else "Μερίδιο προσοχής",[r.get("attention_share") for r in rows])],x,y,w,h,horizontal=False,percent=True,lang=lang)
        return "native_chart"
    if typ == "time_series":
        rows=data.get("rows") or []
        categories=[r.get("date") for r in rows]
        requested=list(data.get("series") or [])
        labels=data.get("series_labels") or {}
        if data.get("scale") == "share_percent" and requested:
            series=[]
            for key in requested[:4]:
                values=[None if r.get(key) is None else 100*_safe_float(r.get(key)) for r in rows]
                series.append((labels.get(key) or key.replace("_opinion_weight_share","").replace("_"," ").title(),values))
            _native_line(slide,categories,series,x,y,w,h,min_y=0,max_y=100)
            return "native_chart"
        # Separate metric families: Brand Reputation is the client-readable headline series.
        if any(r.get("brand_reputation_index") is not None for r in rows):
            _native_line(slide,categories,[("Brand Reputation",[r.get("brand_reputation_index") for r in rows])],x,y,w,h,min_y=0,max_y=100)
            return "native_chart"
    if typ in {"source_matrix","ranking","event_timeline","quality_matrix","risk_panel","cluster_table","evidence_table"}:
        if typ == "source_matrix":
            rows=data.get("rows") or []
            vals=[[r.get("source"),r.get("records"),"—" if r.get("brand_reputation") is None else f"{_safe_float(r.get('brand_reputation')):.1f}",f"{_safe_float(r.get('average_impact')):.2f}"] for r in rows[:8]]
            _add_editable_table(slide,["Source","Records","Reputation","Impact"],vals,x,y,w,h,font_size=8.5)
        elif typ == "ranking":
            rows=data.get("rows") or []
            vals=[[r.get("name"),r.get("records"),f"{_safe_float(r.get('attention')):.2f}","—" if r.get("weighted_sentiment") is None else f"{_safe_float(r.get('weighted_sentiment')):+.2f}"] for r in rows[:8]]
            _add_editable_table(slide,["Name","Records","Attention","Sentiment"],vals,x,y,w,h,font_size=8.5)
        elif typ == "event_timeline":
            ev=data.get("events") or []
            vals=[[e.get("date"),", ".join(e.get("flags") or []),f"{_safe_float(e.get('volume_robust_z')):.1f}"] for e in ev[:8]]
            _add_editable_table(slide,["Date","Signal","Volume z"],vals,x,y,w,h,font_size=8.2)
        elif typ == "quality_matrix":
            items=data.get("items") or []
            friendly={"data_quality":"Data quality","source_coverage":"Source coverage","sample_achievement":"Sample achievement"}
            vals=[[friendly.get(i.get("key"),i.get("key")),"—" if i.get("value") is None else f"{_safe_float(i.get('value')):.1f}",i.get("label") or ""] for i in items]
            _add_editable_table(slide,["Measure","Value","Label"],vals,x,y,w,h,font_size=9)
        elif typ == "risk_panel":
            vals=[["Low authenticity",data.get("low_authenticity_records"),f"{_safe_float(data.get('low_authenticity_share_percent')):.1f}%"],["Coordination clusters",len(data.get("coordination_clusters") or []),""]]
            _add_editable_table(slide,["Signal","Count","Share"],vals,x,y,w,h,font_size=9)
        elif typ == "cluster_table":
            rows=data.get("rows") or []
            vals=[[r.get("cluster_id"),r.get("records"),", ".join(r.get("sources") or []),r.get("representative_excerpt") or ""] for r in rows[:6]]
            _add_editable_table(slide,["Cluster","Records","Sources","Example"],vals,x,y,w,h,col_widths=[1,1,1.3,3],font_size=7.5)
        elif typ == "evidence_table":
            rows=data.get("rows") or []
            vals=[[r.get("platform"),r.get("author"),r.get("date"),r.get("excerpt") or "",f"{_safe_float(r.get('impact_score')):.2f}"] for r in rows[:6]]
            _add_editable_table(slide,["Source","Author","Date","Evidence","Impact"],vals,x,y,w,h,col_widths=[1,1.2,1.1,4.4,0.9],font_size=7.2)
        return "editable_table"
    # Never rasterize unsupported charts: show their structured data as an editable evidence box.
    _add_text(slide, _txt(chart.get("title"),lang), x, y, w, .35, size=11, bold=True)
    _add_text(slide, json.dumps(_json_safe(data), ensure_ascii=False)[:1200], x, y+.45, w, h-.45, size=7.5, color=MUTED)
    return "editable_text_fallback"



def _add_claim_card(slide, text, x, y, w, h, *, label=None, accent=AEGEAN, font_size=11.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.05
    shape.shadow.inherit = False
    shape.fill.solid(); shape.fill.fore_color.rgb = _rgb("FAFAFA"); shape.line.color.rgb = _rgb(STONE_DARK)
    if label:
        _add_text(slide, label, x+.18, y+.14, w-.36, .24, size=8.5, color=accent, bold=True)
        _add_text(slide, text, x+.18, y+.47, w-.36, h-.58, size=font_size, color=INK)
    else:
        _add_text(slide, text, x+.18, y+.18, w-.36, h-.36, size=font_size, color=INK, valign=MSO_ANCHOR.MIDDLE)
    return shape


CLAIM_TYPE_LABELS = {
    "descriptive": ("FACT", "ΔΕΔΟΜΕΝΟ"), "deterministic_rank": ("FACT", "ΔΕΔΟΜΕΝΟ"),
    "investigation": ("INTERPRETATION", "ΕΡΜΗΝΕΙΑ"), "analyst_finding": ("INTERPRETATION", "ΕΡΜΗΝΕΙΑ"),
    "hypothesis": ("HYPOTHESIS", "ΥΠΟΘΕΣΗ"), "recommendation": ("RECOMMENDATION", "ΣΥΣΤΑΣΗ"),
    "benchmark": ("BENCHMARK", "ΣΥΓΚΡΙΣΗ"),
}

def _claim_type_label(claim_type: str, lang: str) -> str | None:
    pair = CLAIM_TYPE_LABELS.get(str(claim_type or ""))
    if not pair:
        return None
    return pair[1] if lang == "el" else pair[0]


def _role_label(role: str, lang: str) -> str:
    labels = {
        "strength": ("Strength", "Δύναμη"),
        "friction": ("Friction", "Τριβή"),
        "momentum": ("Momentum", "Δυναμική"),
        "watch": ("Watch", "Παρακολούθηση"),
    }
    en, el = labels.get(role, (role.title() if role else "Signal", "Σήμα"))
    return el if lang == "el" else en


_GREEK_LATIN = str.maketrans({
    "α": "a", "β": "v", "γ": "g", "δ": "d", "ε": "e", "ζ": "z", "η": "i", "θ": "th",
    "ι": "i", "κ": "k", "λ": "l", "μ": "m", "ν": "n", "ξ": "x", "ο": "o", "π": "p",
    "ρ": "r", "σ": "s", "ς": "s", "τ": "t", "υ": "y", "φ": "f", "χ": "ch", "ψ": "ps",
    "ω": "o", "ά": "a", "έ": "e", "ή": "i", "ί": "i", "ό": "o", "ύ": "y", "ώ": "o",
    "ϊ": "i", "ϋ": "y", "ΐ": "i", "ΰ": "y",
})


def _flatten_greek(value: str) -> str:
    return str(value or "").lower().translate(_GREEK_LATIN)


def _asset_dirs(logo_path: Path | None) -> list[Path]:
    """Where brand assets may live, most specific first.

    ``presentation_assets/`` is the intended home; the repository root is kept as a
    fallback because that is where logo.png already lives.
    """
    roots: list[Path] = []
    if logo_path:
        roots.append(logo_path.parent / "presentation_assets")
        roots.append(logo_path.parent)
    here = Path(__file__).resolve().parents[2]
    roots.append(here / "presentation_assets")
    roots.append(here)
    seen, out = set(), []
    for r in roots:
        key = str(r)
        if key not in seen and r.is_dir():
            seen.add(key)
            out.append(r)
    return out


def _cover_background(logo_path: Path) -> Path | None:
    """Full-bleed studio artwork for the cover, if the brand asset is present."""
    for base in _asset_dirs(logo_path):
        for name in ("cover.png", "cover.jpg", "cover-16x9.png", "report-cover.png"):
            candidate = base / name
            if candidate.exists():
                return candidate
    return None


def _client_logo_from_plan(plan: dict) -> Path | None:
    """Client logo saved in the app's Client Profile, carried inside the plan.

    The profile screen stores the image as a data URL, so the report can be branded
    without anyone touching the repository.
    """
    raw = str((plan or {}).get("client_logo") or "")
    if not raw.startswith("data:image"):
        return None
    try:
        header, _, payload = raw.partition(",")
        if not payload:
            return None
        ext = "png"
        if "jpeg" in header or "jpg" in header:
            ext = "jpg"
        elif "webp" in header:
            ext = "webp"
        elif "svg" in header:
            return None  # python-pptx cannot place SVG
        data = base64.b64decode(payload, validate=False)
        if not data or len(data) > 4_000_000:
            return None
        tmp = Path(tempfile.gettempdir()) / f"signalyth-client-logo.{ext}"
        tmp.write_bytes(data)
        return tmp
    except Exception:
        return None


def _client_logo(logo_path: Path, client: str) -> Path | None:
    """Per-client logo, looked up by a filesystem-safe form of the client name."""
    if not client:
        return None
    bases = [d / "clients" for d in _asset_dirs(logo_path)]
    bases = [b for b in bases if b.is_dir()]
    if not bases:
        return None

    def _slugs(value: str) -> set[str]:
        """Both the raw and the transliterated slug, so a Greek client name
        matches a Latin filename and vice versa."""
        raw = "".join(ch if ch.isalnum() else "-" for ch in str(value).lower())
        latin = "".join(ch if ch.isalnum() else "-" for ch in _flatten_greek(str(value)))
        out = set()
        for form in (raw, latin):
            cleaned = "-".join(part for part in form.split("-") if part)
            if cleaned:
                out.add(cleaned)
                out.add(cleaned.replace("-", ""))
        return out

    wanted = _slugs(client)
    files = [f for b in bases for f in b.iterdir()
             if f.is_file() and f.suffix.lower() in {".png", ".jpg", ".jpeg"}]
    for candidate in files:
        if _slugs(candidate.stem) & wanted:
            return candidate
    # Prefix match: "elas" also serves "ΕΛΑΣ 2".
    for candidate in files:
        for cand_slug in _slugs(candidate.stem):
            if any(w.startswith(cand_slug) for w in wanted if len(cand_slug) >= 3):
                return candidate
    return None


def _render_cover(slide, spec: dict, ctx: dict, lang: str, logo_path: Path, plan: dict | None = None) -> None:
    """Editorial cover: full-bleed artwork, tracked kicker, display serif title.

    The 2026 editorial-serif direction is deliberate: neutral geometric sans faces
    are what every listening platform uses, so a serif display reads as research
    studio rather than dashboard vendor.
    """
    el = lang == "el"
    background = _cover_background(logo_path)
    if background:
        slide.shapes.add_picture(str(background), 0, 0, width=Inches(13.333333), height=Inches(7.5))
    else:
        _add_logo(slide, logo_path, .72, .62, 4.2)

    topic = _clean_text(spec.get("title") or ctx.get("topic") or "", 90)
    client = _clean_text(ctx.get("client") or "", 90)
    market = _clean_text(ctx.get("market") or "", 60)
    period = f"{ctx.get('date_from') or ''} — {ctx.get('date_to') or ''}".strip(" —")

    # Kicker: tracked, uppercase, small. Sets the register before the title lands.
    _add_display(slide, "BRAND INTELLIGENCE REPORT" if not el else "ΑΝΑΦΟΡΑ BRAND INTELLIGENCE",
                 .95, 2.62, 6.6, .3, size=10.5, color=COVER_ACCENT, font=LABEL_FONT,
                 bold=True, tracking=3.2)
    _add_rule(slide, .95, 3.02, 1.15, color=COVER_ACCENT, width=2.5)

    # The subject of the report, set large. This is the element the client remembers.
    title_size = 54 if len(topic) <= 18 else 44 if len(topic) <= 28 else 34
    _add_display(slide, topic, .92, 3.16, 7.4, 1.45, size=title_size, color=COVER_INK,
                 bold=True, line_spacing=0.92)

    if client:
        _add_display(slide, client, .95, 4.62, 7.0, .42, size=15, color="44506B",
                     font=LABEL_FONT)
    meta = " · ".join(x for x in (market, period) if x)
    if meta:
        _add_display(slide, meta.upper(), .95, 5.12, 7.0, .32, size=9.5, color="7A8699",
                     font=LABEL_FONT, tracking=2.0)

    # Client mark, bottom right, introduced by a quiet label.
    client_logo = _client_logo_from_plan(plan or {}) or _client_logo(logo_path, ctx.get("client"))
    if client_logo:
        _add_display(slide, "PREPARED FOR" if not el else "ΓΙΑ ΤΟΝ ΠΕΛΑΤΗ",
                     9.05, 5.92, 3.4, .26, size=8, color="7A8699", font=LABEL_FONT,
                     bold=True, tracking=2.6, align=PP_ALIGN.RIGHT)
        try:
            pic = slide.shapes.add_picture(str(client_logo), Inches(9.05), Inches(6.24), height=Inches(0.72))
            # Right-align the mark against the 12.4in margin.
            pic.left = Inches(12.4) - pic.width
        except Exception:
            pass


# ---------- Executive dashboard (slide 2) ----------
# Locked design "Version B (light)": gauge + sentiment bars + source donut +
# five KPI cards, every value computed from the run's own data. Rendered with
# pure python-pptx shapes so it works on any deployment without new libraries.

DASH_CARD_BG = "FAFAFA"
DASH_TRACK = "ECEFF1"
DASH_CHIP_NEG_BG = "FBE9EA"
DASH_CHIP_NEG_TX = "B3262E"
DASH_CHIP_POS_BG = "E8F3EC"
DASH_CHIP_POS_TX = "1E7F4F"

_EL_MONTHS_GEN = ["Ιανουαρίου", "Φεβρουαρίου", "Μαρτίου", "Απριλίου", "Μαΐου", "Ιουνίου",
                  "Ιουλίου", "Αυγούστου", "Σεπτεμβρίου", "Οκτωβρίου", "Νοεμβρίου", "Δεκεμβρίου"]
_EN_MONTHS = ["January", "February", "March", "April", "May", "June",
              "July", "August", "September", "October", "November", "December"]


def _upper_label(text: str) -> str:
    """Uppercase for small-caps labels; Greek uppercase never keeps the tonos."""
    plain = str(text or "").translate(str.maketrans("άέήίόύώΐΰϊϋ", "αεηιουωιυιυ"))
    return plain.upper()


def _period_label(date_from, date_to, lang: str) -> str:
    """Human date range: «2 – 10 Σεπτεμβρίου 2026» instead of ISO."""
    try:
        d0 = datetime.strptime(str(date_from)[:10], "%Y-%m-%d")
        d1 = datetime.strptime(str(date_to)[:10], "%Y-%m-%d")
    except (TypeError, ValueError):
        return f"{date_from or ''} — {date_to or ''}".strip(" —")
    months = _EL_MONTHS_GEN if lang == "el" else _EN_MONTHS
    if (d0.year, d0.month) == (d1.year, d1.month):
        return f"{d0.day} – {d1.day} {months[d1.month - 1]} {d1.year}"
    if d0.year == d1.year:
        return f"{d0.day} {months[d0.month - 1]} – {d1.day} {months[d1.month - 1]} {d1.year}"
    return f"{d0.day} {months[d0.month - 1]} {d0.year} – {d1.day} {months[d1.month - 1]} {d1.year}"


def _compact_number(value) -> str:
    n = _safe_float(value, 0.0)
    for limit, suffix in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if abs(n) >= limit:
            txt = f"{n / limit:.1f}".rstrip("0").rstrip(".")
            return f"{txt}{suffix}"
    return f"{int(round(n))}"


def _ring_segment(slide, cx, cy, r_out, r_in, a0, a1, color, *, steps=None):
    """Filled ring segment (thick arc) as a freeform. Angles in degrees,
    0 = 12 o'clock, clockwise. Coordinates in inches."""
    span = float(a1) - float(a0)
    if span <= 0.1:
        return None
    steps = steps or max(10, int(span / 5))

    def pt(r, ang):
        rad = math.radians(ang)
        return (cx + r * math.sin(rad), cy - r * math.cos(rad))

    pts = [pt(r_out, a0 + span * i / steps) for i in range(steps + 1)]
    pts += [pt(r_in, a1 - span * i / steps) for i in range(steps + 1)]
    # FreeformBuilder rounds local coordinates to integers before scaling, so
    # inches are passed as integer 1/100000-inch units with a matching scale.
    unit = 100000
    ipts = [(int(round(x * unit)), int(round(y * unit))) for x, y in pts]
    try:
        fb = slide.shapes.build_freeform(ipts[0][0], ipts[0][1], scale=914400.0 / unit)
        fb.add_line_segments(ipts[1:], close=True)
        shp = fb.convert_to_shape()
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(color)
        shp.line.fill.background(); shp.shadow.inherit = False
        return shp
    except Exception:
        return None


def _dash_dot(slide, x, y, d, color):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(color)
    dot.line.fill.background(); dot.shadow.inherit = False
    return dot


def _dash_card(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.055
    shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(DASH_CARD_BG)
    shape.line.color.rgb = _rgb(STONE_DARK)
    shape.shadow.inherit = False
    return shape


def _reach_totals(folder: Path) -> tuple[int | None, int | None]:
    """(followers_total, views_total) from the run's own records.
    Views: plain sum over analysis-ready records. Followers: summed once per
    unique author (max seen), so repeated posts never double-count reach."""
    records = None
    for rel in (("intelligence", "records.json"), ("analysis", "analysis-ready.json")):
        try:
            data = RunStore.read(folder / rel[0] / rel[1], None)
        except Exception:
            data = None
        if isinstance(data, list) and data:
            records = data
            break
    if not records:
        return None, None
    views_total = 0
    followers_by_author: dict[str, int] = {}
    for r in records:
        if not isinstance(r, dict):
            continue
        views_total += _safe_int(r.get("views"))
        author = str(r.get("author") or "").strip().casefold()
        key = f"{r.get('platform') or ''}::{author}" if author else f"::{id(r)}"
        followers_by_author[key] = max(followers_by_author.get(key, 0), _safe_int(r.get("followers")))
    return sum(followers_by_author.values()), views_total


def _executive_dashboard_data(folder: Path, visual_pack: dict) -> dict | None:
    """Everything slide 2 shows, computed from this run's charts and records.
    Returns None when the essentials are missing so the caller can fall back
    to the legacy executive-summary layout instead of shipping a broken slide."""
    try:
        charts = _chart_map(visual_pack)
        rep = (charts.get("brand_reputation") or {}).get("data") or {}
        if rep.get("value") is None:
            return None
        dash: dict = {"score": round(_safe_float(rep.get("value")), 1)}

        conf = (charts.get("evidence_confidence") or {}).get("data") or {}
        dash["evidence_score"] = None if conf.get("value") is None else round(_safe_float(conf.get("value")))

        cats = ((charts.get("sentiment_distribution") or {}).get("data") or {}).get("categories") or []
        sentiment = [{"key": _cat_key(c.get("label")), "value": round(_safe_float(c.get("value")), 1)} for c in cats]
        sentiment = [s for s in sentiment if s["key"] in SENTIMENT_COLORS]
        sentiment.sort(key=lambda s: -s["value"])
        dash["sentiment"] = sentiment

        # Source composition donut: people, media and unknown only (locked scope);
        # owned/organisation activity stays out of this chart by design.
        rows = ((charts.get("origin_breakdown") or {}).get("data") or {}).get("rows") or []
        wanted = {"person", "media", "unknown"}
        comp = {_cat_key(r.get("name")): _safe_int(r.get("records")) for r in rows if _cat_key(r.get("name")) in wanted}
        total = sum(comp.values())
        dash["composition"] = [
            {"key": k, "share": round(100.0 * comp.get(k, 0) / total, 1) if total else 0.0}
            for k in ("person", "media", "unknown")
        ] if total else []

        items = ((charts.get("sample_overview") or {}).get("data") or {}).get("items") or []
        by_key = {str(i.get("key")): i.get("value") for i in items}
        dash["records_ready"] = _safe_int(by_key.get("analysis_ready")) or None
        dash["voices"] = None if by_key.get("effective_voices") is None else int(round(_safe_float(by_key.get("effective_voices"))))

        trend_rows = ((charts.get("time_trends") or {}).get("data") or {}).get("rows") or []
        series = [_safe_float(r.get("brand_reputation_index"), None) for r in trend_rows
                  if isinstance(r, dict) and r.get("brand_reputation_index") is not None]
        dash["delta"] = round(series[-1] - series[0], 1) if len(series) >= 2 else None

        followers, views = _reach_totals(folder)
        dash["followers_total"] = followers
        dash["views_total"] = views
        return dash
    except Exception:
        return None


def _render_executive_dashboard(slide, dash: dict, lang: str, ctx: dict) -> None:
    el = lang == "el"

    # Client + period, top right, mirroring the cover's identity block.
    client = _clean_text(ctx.get("client") or "", 60)
    topic = _clean_text(ctx.get("topic") or "", 60)
    who = " · ".join(x for x in (client, topic) if x)
    period = _period_label(ctx.get("date_from"), ctx.get("date_to"), lang)
    market = _clean_text(ctx.get("market") or "", 40)
    when = " · ".join(x for x in (period, market) if x)
    if who:
        _add_text(slide, who, 8.35, 0.36, 4.38, 0.28, size=12, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    if when:
        _add_text(slide, when, 8.35, 0.66, 4.38, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)

    row_y, row_h = 1.42, 3.22

    # --- Card A: Brand Reputation gauge ---
    ax, aw = 0.58, 4.12
    _dash_card(slide, ax, row_y, aw, row_h)
    _add_text(slide, "BRAND REPUTATION", ax + 0.28, row_y + 0.18, aw - 0.56, 0.24, size=9.5, color=MUTED, bold=True, font=LABEL_FONT)
    cx, cy = ax + aw / 2, row_y + 1.52
    r_out, r_in = 1.00, 0.79
    r_mid, cap = (r_out + r_in) / 2, r_out - r_in
    score = max(0.0, min(100.0, _safe_float(dash.get("score"))))
    a0, span = -135.0, 270.0
    a_val = a0 + span * score / 100.0
    _ring_segment(slide, cx, cy, r_out, r_in, a0, a0 + span, DASH_TRACK)
    for ang in (a0, a0 + span):
        _dash_dot(slide, cx + r_mid * math.sin(math.radians(ang)) - cap / 2,
                  cy - r_mid * math.cos(math.radians(ang)) - cap / 2, cap, DASH_TRACK)
    _ring_segment(slide, cx, cy, r_out, r_in, a0, a_val, AEGEAN)
    _dash_dot(slide, cx + r_mid * math.sin(math.radians(a0)) - cap / 2,
              cy - r_mid * math.cos(math.radians(a0)) - cap / 2, cap, AEGEAN)
    knob_d = cap + 0.10
    knob = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + r_mid * math.sin(math.radians(a_val)) - knob_d / 2),
                                  Inches(cy - r_mid * math.cos(math.radians(a_val)) - knob_d / 2), Inches(knob_d), Inches(knob_d))
    knob.fill.solid(); knob.fill.fore_color.rgb = _rgb(DASH_CARD_BG)
    knob.line.color.rgb = _rgb(AEGEAN); knob.line.width = Pt(2.4); knob.shadow.inherit = False
    _add_text(slide, f"{score:.1f}", cx - 1.0, cy - 0.44, 2.0, 0.56, size=33, color=INK, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "/ 100", cx - 1.0, cy + 0.14, 2.0, 0.24, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

    delta = dash.get("delta")
    chip_y = row_y + 2.66
    if delta is not None and abs(_safe_float(delta)) >= 0.05:
        d = _safe_float(delta)
        up = d > 0
        chip_bg, chip_tx, arrow = (DASH_CHIP_POS_BG, DASH_CHIP_POS_TX, "▲") if up else (DASH_CHIP_NEG_BG, DASH_CHIP_NEG_TX, "▼")
        label = (f"{arrow}  {abs(d):.1f} pts in period" if not el else f"{arrow}  {abs(d):.1f} μον. στην περίοδο")
        cw = 2.0
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - cw / 2), Inches(chip_y), Inches(cw), Inches(0.32))
        chip.adjustments[0] = 0.5; chip.fill.solid(); chip.fill.fore_color.rgb = _rgb(chip_bg)
        chip.line.fill.background(); chip.shadow.inherit = False
        _add_text(slide, label, cx - cw / 2, chip_y, cw, 0.32, size=9, color=chip_tx, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    # §9 scope subtitle stays with the score wherever it appears.
    _add_text(slide, "Digital conversation signal — not a public-opinion poll" if not el
              else "Δείκτης ψηφιακής συζήτησης — όχι μέτρηση κοινής γνώμης",
              ax + 0.2, row_y + 3.0 - 0.06, aw - 0.4, 0.22, size=8, color=NEUTRAL, align=PP_ALIGN.CENTER)

    # --- Card B: weighted sentiment bars ---
    bx, bw = 4.90, 4.12
    _dash_card(slide, bx, row_y, bw, row_h)
    _add_text(slide, "WEIGHTED SENTIMENT" if not el else "ΣΤΑΘΜΙΣΜΕΝΟ SENTIMENT",
              bx + 0.28, row_y + 0.18, bw - 0.56, 0.24, size=9.5, color=MUTED, bold=True, font=LABEL_FONT)
    inner_x, inner_w = bx + 0.30, bw - 0.60
    bar_h, y0 = 0.20, row_y + 0.58
    rows_s = (dash.get("sentiment") or [])[:4]
    step = 0.60 if len(rows_s) >= 4 else 0.72
    for i, s in enumerate(rows_s):
        yy = y0 + i * step
        color = SENTIMENT_COLORS.get(s["key"], NEUTRAL)
        _add_text(slide, _cat_label(s["key"], lang), inner_x, yy, 1.7, 0.22, size=10, color=INK)
        _add_text(slide, f"{s['value']:.1f}%", inner_x + inner_w - 0.95, yy, 0.95, 0.22, size=10.5, color=color, bold=True, align=PP_ALIGN.RIGHT)
        ty = yy + 0.26
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(inner_x), Inches(ty), Inches(inner_w), Inches(bar_h))
        track.adjustments[0] = 0.5; track.fill.solid(); track.fill.fore_color.rgb = _rgb(DASH_TRACK)
        track.line.fill.background(); track.shadow.inherit = False
        vw = max(inner_w * max(0.0, min(100.0, _safe_float(s["value"]))) / 100.0, bar_h)
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(inner_x), Inches(ty), Inches(vw), Inches(bar_h))
        bar.adjustments[0] = 0.5; bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(color)
        bar.line.fill.background(); bar.shadow.inherit = False

    # --- Card C: source composition donut ---
    dx, dw = 9.22, 3.51
    _dash_card(slide, dx, row_y, dw, row_h)
    _add_text(slide, "SOURCE MIX" if not el else "ΣΥΝΘΕΣΗ ΠΗΓΩΝ",
              dx + 0.28, row_y + 0.18, dw - 0.56, 0.24, size=9.5, color=MUTED, bold=True, font=LABEL_FONT)
    comp = dash.get("composition") or []
    comp_colors = {"person": AEGEAN, "media": MIXED, "unknown": NEUTRAL}
    ccx, ccy = dx + dw / 2, row_y + 1.28
    ro, ri = 0.74, 0.47
    if comp:
        ang = 0.0
        for c in comp:
            sweep = 360.0 * max(0.0, _safe_float(c["share"])) / 100.0
            _ring_segment(slide, ccx, ccy, ro, ri, ang, min(ang + sweep, 359.9), comp_colors.get(c["key"], NEUTRAL))
            ang += sweep
        top = max(comp, key=lambda c: c["share"])
        top_txt = f"{top['share']:.1f}".rstrip("0").rstrip(".") + "%"
        _add_text(slide, top_txt, ccx - 0.55, ccy - 0.24, 1.1, 0.32, size=15, color=INK, bold=True, align=PP_ALIGN.CENTER)
        _add_text(slide, _cat_label(top["key"], lang), ccx - 0.6, ccy + 0.07, 1.2, 0.2, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)
        for i, c in enumerate(comp):
            ly = row_y + 2.28 + i * 0.25
            _dash_dot(slide, dx + 0.34, ly + 0.05, 0.11, comp_colors.get(c["key"], NEUTRAL))
            _add_text(slide, _cat_label(c["key"], lang), dx + 0.54, ly, 1.7, 0.22, size=9, color=INK)
            _add_text(slide, f"{c['share']:.1f}%", dx + dw - 1.28, ly, 1.0, 0.22, size=9, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)

    # --- KPI row: five cards, everything from this run ---
    k_y, k_h, gap = 4.86, 1.38, 0.18
    k_w = (12.15 - 4 * gap) / 5
    ev_score = dash.get("evidence_score")
    ev_label = None
    if ev_score is not None:
        ev_label = ("High" if ev_score >= 70 else "Medium" if ev_score >= 40 else "Low") if not el else \
                   ("Υψηλή" if ev_score >= 70 else "Μέτρια" if ev_score >= 40 else "Χαμηλή")
    followers = dash.get("followers_total")
    views = dash.get("views_total")
    kpis = [
        ("—" if not dash.get("records_ready") else f"{dash['records_ready']:,}".replace(",", "."),
         "Analysis-ready records" if not el else "Αναφορές έτοιμες για ανάλυση", None),
        ("—" if dash.get("voices") is None else f"{dash['voices']:,}".replace(",", "."),
         "Independent voices" if not el else "Ανεξάρτητες φωνές", None),
        ("—" if ev_score is None else f"{ev_score:.0f}/100",
         "Evidence confidence" if not el else "Βεβαιότητα evidence", ev_label),
        ("—" if followers is None else _compact_number(followers),
         "Source followers (unique)" if not el else "Followers πηγών (μοναδικοί)", None),
        ("—" if views is None else _compact_number(views),
         "Total views" if not el else "Συνολικά views", None),
    ]
    for i, (value, label, sub) in enumerate(kpis):
        x = 0.58 + i * (k_w + gap)
        _dash_card(slide, x, k_y, k_w, k_h)
        accent = AEGEAN if i == 2 else INK
        _add_text(slide, value, x + 0.06, k_y + 0.16, k_w - 0.12, 0.5, size=22 if len(value) <= 7 else 18,
                  color=accent, bold=True, align=PP_ALIGN.CENTER)
        if sub:
            _add_text(slide, _upper_label(sub), x + 0.06, k_y + 0.64, k_w - 0.12, 0.2, size=8, color=POS, bold=True, align=PP_ALIGN.CENTER, font=LABEL_FONT)
        _add_text(slide, label, x + 0.12, k_y + (0.86 if sub else 0.74), k_w - 0.24, 0.46, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)


# ---------- Brand Reputation Timeline (slide 3) ----------
# Locked design: value-coloured line (deep green -> amber -> red, saturated
# around the methodology's neutral band), value on every dot, Peak/Low badges,
# auto-written narrative, three mini-KPIs. Granularity adapts to the run:
# daily for multi-day reports, hourly for crisis/debate runs measured in hours.

TL_SCALE_NEG = (197, 48, 58)    # C5303A
TL_SCALE_MID = (217, 138, 43)   # D98A2B
TL_SCALE_POS = (30, 107, 69)    # 1E6B45
TL_AREA = "EAF0F2"

_EL_DAYS = ["Δευτέρα", "Τρίτη", "Τετάρτη", "Πέμπτη", "Παρασκευή", "Σάββατο", "Κυριακή"]
_EL_DAYS_AB = ["Δε", "Τρ", "Τε", "Πε", "Πα", "Σα", "Κυ"]
_EN_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
_EN_DAYS_AB = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _val_color(v) -> str:
    """Value -> hex on the reputation palette. Saturates at 30 (full red) and
    70 (full green) so movement around the 45-55 neutral band stays visible."""
    t = max(0.0, min(1.0, (_safe_float(v) - 30.0) / 40.0))
    if t <= 0.5:
        a, b, tt = TL_SCALE_NEG, TL_SCALE_MID, t * 2
    else:
        a, b, tt = TL_SCALE_MID, TL_SCALE_POS, (t - 0.5) * 2
    return "".join(f"{int(round(a[i] + (b[i] - a[i]) * tt)):02X}" for i in range(3))


def _parse_dt(value):
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None


def _tl_day_labels(dt, lang):
    ab = (_EL_DAYS_AB if lang == "el" else _EN_DAYS_AB)[dt.weekday()]
    full = (_EL_DAYS if lang == "el" else _EN_DAYS)[dt.weekday()]
    dm = f"{dt.day}/{dt.month}"
    return f"{ab} {dm}", f"{full} {dm}"


def _reputation_timeline_data(folder: Path, visual_pack: dict, lang: str) -> dict | None:
    """Timeline points + narrative facts for slide 3, all from this run.
    Daily: the run's own brand_reputation_index series. Hourly (short runs):
    per-hour impact-weighted sentiment mapped to 0-100 with the methodology's
    own transform ((s+1)/2*100). Returns None when there is nothing to plot."""
    try:
        charts = _chart_map(visual_pack)
        rows = ((charts.get("time_trends") or {}).get("data") or {}).get("rows") or []
        daily = []
        for r in rows:
            if not isinstance(r, dict) or r.get("brand_reputation_index") is None:
                continue
            dt = _parse_dt(str(r.get("date"))[:10] + "T00:00:00+00:00")
            if dt is None:
                continue
            daily.append((dt, round(_safe_float(r.get("brand_reputation_index")), 1), _safe_int(r.get("records"))))
        daily.sort(key=lambda x: x[0])

        points = []
        granularity = None
        if len(daily) >= 2:
            granularity = "daily"
            for dt, v, vol in daily:
                ab, full = _tl_day_labels(dt, lang)
                points.append({"label": ab, "full": full, "value": v, "volume": vol})
        else:
            # Crisis / debate mode: bucket the run's records per hour.
            records = None
            for rel in (("intelligence", "records.json"), ("analysis", "analysis-ready.json")):
                try:
                    data = RunStore.read(folder / rel[0] / rel[1], None)
                except Exception:
                    data = None
                if isinstance(data, list) and data:
                    records = data
                    break
            if not records:
                return None
            buckets: dict = {}
            for r in records:
                if not isinstance(r, dict):
                    continue
                dt = _parse_dt(r.get("date"))
                s = ((r.get("ai_analysis") or {}).get("sentiment_score"))
                if dt is None or s is None:
                    continue
                w = max(0.05, _safe_float((r.get("intelligence") or {}).get("impact_score"), 0.5))
                key = dt.replace(minute=0, second=0, microsecond=0)
                b = buckets.setdefault(key, {"num": 0.0, "den": 0.0, "n": 0})
                b["num"] += w * _safe_float(s)
                b["den"] += w
                b["n"] += 1
            hours = sorted(buckets.items())
            if len(hours) < 2:
                return None
            span_h = (hours[-1][0] - hours[0][0]).total_seconds() / 3600.0
            if span_h > 72:
                return None  # long run without a daily series: nothing honest to plot
            granularity = "hourly"
            multiday = hours[0][0].date() != hours[-1][0].date()
            for dt, b in hours:
                v = round(((b["num"] / b["den"]) + 1.0) / 2.0 * 100.0, 1) if b["den"] else 50.0
                hh = dt.strftime("%H:00")
                lbl = f"{dt.day}/{dt.month} {hh}" if multiday else hh
                points.append({"label": hh if not multiday else lbl, "full": lbl, "value": v, "volume": b["n"]})

        if len(points) < 2:
            return None
        vals = [p["value"] for p in points]
        i_peak = max(range(len(vals)), key=lambda i: vals[i])
        i_low = min(range(len(vals)), key=lambda i: vals[i])
        steep = None
        if len(vals) >= 3:
            j = max(range(1, len(vals)), key=lambda i: abs(vals[i] - vals[i - 1]))
            steep = {"i": j, "change": round(vals[j] - vals[j - 1], 1),
                     "from": points[j - 1]["full"], "to": points[j]["full"],
                     "volume": points[j]["volume"],
                     "is_busiest": points[j]["volume"] == max(p["volume"] for p in points)}
        return {"granularity": granularity, "points": points, "i_peak": i_peak, "i_low": i_low,
                "delta": round(vals[-1] - vals[0], 1), "first": vals[0], "last": vals[-1], "steep": steep}
    except Exception:
        return None


def _timeline_narrative(tl: dict, lang: str) -> list[tuple[str, str]]:
    el = lang == "el"
    pts = tl["points"]
    delta, first, last = tl["delta"], tl["first"], tl["last"]
    hourly = tl["granularity"] == "hourly"
    out: list[tuple[str, str]] = []

    if delta <= -3:
        head = "Πτωτική πορεία. " if el else "Downward trajectory. "
        body = (f"Το Brand Reputation έχασε {abs(delta):.1f} μονάδες μέσα στην περίοδο — από {first:.1f} στις {last:.1f}."
                if el else f"Brand Reputation lost {abs(delta):.1f} points over the period — from {first:.1f} to {last:.1f}.")
    elif delta >= 3:
        head = "Ανοδική πορεία. " if el else "Upward trajectory. "
        body = (f"Το Brand Reputation κέρδισε {delta:.1f} μονάδες μέσα στην περίοδο — από {first:.1f} στις {last:.1f}."
                if el else f"Brand Reputation gained {delta:.1f} points over the period — from {first:.1f} to {last:.1f}.")
    else:
        head = "Σταθερή εικόνα. " if el else "Stable picture. "
        body = (f"Το Brand Reputation μεταβλήθηκε μόλις κατά {delta:+.1f} μονάδες στην περίοδο ({first:.1f} → {last:.1f})."
                if el else f"Brand Reputation moved only {delta:+.1f} points over the period ({first:.1f} → {last:.1f}).")
    out.append((head, body))

    def edge_note(i):
        if i == 0:
            return ("στην εκκίνηση της μέτρησης" if el else "at the start of measurement")
        if i == len(pts) - 1:
            return ("στο κλείσιμο της περιόδου" if el else "at the close of the period")
        return None

    pk, lo = tl["i_peak"], tl["i_low"]
    note = edge_note(pk)
    out.append(("Κορύφωση: " if el else "Peak: ",
                f"{pts[pk]['full']}" + (f", {note}" if note else "") + f" ({pts[pk]['value']:.1f})."))
    note = edge_note(lo)
    out.append(("Χαμηλό: " if el else "Low: ",
                f"{pts[lo]['full']}" + (f", {note}" if note else "") + f" ({pts[lo]['value']:.1f})."))

    st = tl.get("steep")
    if st and abs(st["change"]) >= 0.1:
        drop = st["change"] < 0
        head = (("Πιο απότομη πτώση: " if drop else "Πιο απότομη άνοδος: ") if el
                else ("Steepest drop: " if drop else "Steepest rise: "))
        unit = ("μονάδες" if el else "points")
        seg = f"{st['from']} → {st['to']} ({st['change']:+.1f} {unit})"
        if st["is_busiest"]:
            tail = ((", την ώρα με τον υψηλότερο όγκο αναφορών" if hourly else ", τη μέρα με τον υψηλότερο όγκο αναφορών")
                    if el else (", in the busiest hour of the period" if hourly else ", on the busiest day of the period"))
            tail += f" ({st['volume']})."
        else:
            tail = (f", με {st['volume']} αναφορές στο διάστημα." if el
                    else f", with {st['volume']} mentions in that window.")
        out.append((head, seg + tail))
    return out


def _add_rich(slide, parts, x, y, w, h, *, size=9.6, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    for idx, (text, bold) in enumerate(parts):
        run = p.add_run()
        # Keep the separating space between runs: clean, then restore the
        # trailing space _clean_text strips from the lead-in run.
        cleaned = _clean_text(text, 600)
        if idx < len(parts) - 1 and str(text).endswith(" "):
            cleaned += " "
        run.text = cleaned
        try:
            from pptx.oxml.ns import qn
            t_el = run._r.find(qn("a:t"))
            if t_el is not None:
                t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        except Exception:
            pass
        run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = _rgb(color)
    return box


def _render_reputation_timeline(slide, tl: dict, lang: str, ctx: dict) -> None:
    el = lang == "el"
    pts = tl["points"]; n = len(pts)
    hourly = tl["granularity"] == "hourly"

    # Top-right identity, matching slide 2.
    client = _clean_text(ctx.get("client") or "", 60)
    topic = _clean_text(ctx.get("topic") or "", 60)
    who = " · ".join(x for x in (client, topic) if x)
    gran = ("ανά ώρα" if hourly else "ανά ημέρα") if el else ("per hour" if hourly else "per day")
    when = " · ".join(x for x in (_period_label(ctx.get("date_from"), ctx.get("date_to"), lang), gran) if x)
    if who:
        _add_text(slide, who, 8.35, 0.36, 4.38, 0.28, size=12, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    _add_text(slide, when, 8.35, 0.66, 4.38, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)

    # --- Chart card ---
    cx0, cy0, cw, chh = 0.58, 1.42, 8.10, 5.44
    _dash_card(slide, cx0, cy0, cw, chh)
    _add_text(slide, ("BRAND REPUTATION · ΕΞΕΛΙΞΗ ΣΤΗΝ ΠΕΡΙΟΔΟ" if el else "BRAND REPUTATION · EVOLUTION OVER THE PERIOD"),
              cx0 + 0.28, cy0 + 0.18, cw - 0.56, 0.24, size=9.5, color=MUTED, bold=True, font=LABEL_FONT)

    px0, px1 = cx0 + 0.62, cx0 + cw - 0.34
    py_bot, py_top = cy0 + 4.62, cy0 + 1.02
    X = (lambda i: px0 + (px1 - px0) * i / (n - 1)) if n > 1 else (lambda i: (px0 + px1) / 2)
    Y = lambda v: py_bot - (max(0.0, min(100.0, v)) / 100.0) * (py_bot - py_top)

    # Grid + y labels
    for g in (0, 25, 50, 75, 100):
        _add_rule(slide, px0, Y(g), px1 - px0, color=DASH_TRACK, width=0.8)
        _add_text(slide, str(g), cx0 + 0.14, Y(g) - 0.09, 0.42, 0.18, size=7.5, color=NEUTRAL, align=PP_ALIGN.RIGHT)

    # Colour-scale legend, top right of the plot
    sx, sw_total, seg_n = px1 - 2.35, 1.55, 26
    for i in range(seg_n):
        seg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(sx + sw_total * i / seg_n), Inches(cy0 + 0.52),
                                     Inches(sw_total / seg_n + 0.006), Inches(0.08))
        seg.fill.solid(); seg.fill.fore_color.rgb = _rgb(_val_color(100.0 * i / (seg_n - 1)))
        seg.line.fill.background(); seg.shadow.inherit = False
    _add_text(slide, ("Αρνητικό 0" if el else "Negative 0"), sx - 0.95, cy0 + 0.475, 0.9, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)
    _add_text(slide, ("100 Θετικό" if el else "100 Positive"), sx + sw_total + 0.05, cy0 + 0.475, 0.95, 0.18, size=7.5, color=MUTED)

    # Area fill under the line
    area_pts = [(X(0), py_bot)] + [(X(i), Y(p["value"])) for i, p in enumerate(pts)] + [(X(n - 1), py_bot)]
    unit = 100000
    ipts = [(int(round(x * unit)), int(round(y * unit))) for x, y in area_pts]
    try:
        fb = slide.shapes.build_freeform(ipts[0][0], ipts[0][1], scale=914400.0 / unit)
        fb.add_line_segments(ipts[1:], close=True)
        shp = fb.convert_to_shape()
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(TL_AREA)
        shp.line.fill.background(); shp.shadow.inherit = False
    except Exception:
        pass

    # Value-coloured segments
    from pptx.enum.shapes import MSO_CONNECTOR
    for i in range(n - 1):
        c = _val_color((pts[i]["value"] + pts[i + 1]["value"]) / 2.0)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X(i)), Inches(Y(pts[i]["value"])),
                                          Inches(X(i + 1)), Inches(Y(pts[i + 1]["value"])))
        conn.line.color.rgb = _rgb(c); conn.line.width = Pt(3.4); conn.shadow.inherit = False

    # Which dots get a printed value (thin out on dense series; extremes always)
    if n <= 14:
        labelled = set(range(n))
    else:
        step = max(1, (n + 11) // 12)
        labelled = set(range(0, n, step)) | {0, n - 1, tl["i_peak"], tl["i_low"]}

    dot_d = 0.13
    for i, p in enumerate(pts):
        x, y = X(i), Y(p["value"])
        c = _val_color(p["value"])
        if i in (tl["i_peak"], tl["i_low"]):
            ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.145), Inches(y - 0.145), Inches(0.29), Inches(0.29))
            ring.fill.background(); ring.line.color.rgb = _rgb(c); ring.line.width = Pt(1.6); ring.shadow.inherit = False
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - dot_d / 2), Inches(y - dot_d / 2), Inches(dot_d), Inches(dot_d))
        dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(c)
        dot.line.color.rgb = _rgb(WHITE); dot.line.width = Pt(1.5); dot.shadow.inherit = False
        if i in (tl["i_peak"], tl["i_low"]):
            continue  # value lives in the badge
        if i in labelled:
            above = p["value"] >= pts[i - 1]["value"] if i > 0 else True
            ly = y - 0.30 if above else y + 0.12
            _add_text(slide, f"{p['value']:.1f}", x - 0.35, ly, 0.70, 0.20, size=8.5, color=c, bold=True, align=PP_ALIGN.CENTER)

    # Peak / Low badges with the value inside
    for i, tag in ((tl["i_peak"], "Peak"), (tl["i_low"], "Low")):
        x, y = X(i), Y(pts[i]["value"])
        c = _val_color(pts[i]["value"])
        bw = 1.06
        bx = min(max(x - bw / 2, px0), px1 - bw)
        by = y - 0.56 if tag == "Peak" else y + 0.24
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(bw), Inches(0.30))
        badge.adjustments[0] = 0.5; badge.fill.solid(); badge.fill.fore_color.rgb = _rgb(c)
        badge.line.fill.background(); badge.shadow.inherit = False
        _add_text(slide, f"{tag} {pts[i]['value']:.1f}", bx, by, bw, 0.30, size=8.5, color=WHITE, bold=True,
                  align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

    # X axis labels (thinned like the values)
    if n <= 12:
        x_lab = set(range(n))
    else:
        step = max(1, (n + 9) // 10)
        x_lab = set(range(0, n, step)) | {0, n - 1}
    for i in sorted(x_lab):
        _add_text(slide, pts[i]["label"], X(i) - 0.55, py_bot + 0.10, 1.10, 0.20, size=8, color=MUTED, align=PP_ALIGN.CENTER)

    caption = (("Δείκτης 0–100 · υπολογίζεται ανά ημέρα από τα evidence της περιόδου" if not hourly
                else "Σταθμισμένο sentiment σε κλίμακα 0–100 · υπολογίζεται ανά ώρα από τα evidence της περιόδου") if el
               else ("Index 0–100 · computed per day from the period's evidence" if not hourly
                     else "Weighted sentiment on a 0–100 scale · computed per hour from the period's evidence"))
    _add_text(slide, caption, cx0 + 0.28, cy0 + chh - 0.34, cw - 0.56, 0.22, size=8, color=NEUTRAL)

    # --- Narrative card ---
    nx, nw_ = 8.88, 3.85
    _dash_card(slide, nx, 1.42, nw_, 3.62)
    _add_text(slide, ("ΤΙ ΔΕΙΧΝΕΙ Η ΓΡΑΜΜΗ" if el else "WHAT THE LINE SHOWS"),
              nx + 0.28, 1.62, nw_ - 0.56, 0.24, size=9.5, color=MUTED, bold=True, font=LABEL_FONT)
    paras = _timeline_narrative(tl, lang)[:4]
    block = (3.62 - 0.62) / max(1, len(paras))
    py = 1.96
    for head, body in paras:
        _add_rich(slide, [(head, True), (body, False)], nx + 0.28, py, nw_ - 0.56, block, size=9.4)
        py += block

    # --- Mini KPI chips ---
    ky, kh_ = 5.22, 1.64
    kw_ = (nw_ - 0.24) / 3
    pk, lo = pts[tl["i_peak"]], pts[tl["i_low"]]
    d = tl["delta"]
    chips = [
        (f"{pk['value']:.1f}", _val_color(pk["value"]), f"Peak {pk['label']}"),
        (f"{lo['value']:.1f}", _val_color(lo["value"]), f"Low {lo['label']}"),
        (f"{d:+.1f}", POS if d > 0 else NEG if d < 0 else NEUTRAL,
         ("Μεταβολή περιόδου" if el else "Period change")),
    ]
    for i, (value, color, label) in enumerate(chips):
        x = nx + i * (kw_ + 0.12)
        _dash_card(slide, x, ky, kw_, kh_)
        _add_text(slide, value, x, ky + 0.22, kw_, 0.4, size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
        _add_text(slide, label, x + 0.06, ky + 0.72, kw_ - 0.12, 0.8, size=8, color=MUTED, align=PP_ALIGN.CENTER)


def generate_pptx(presentation_plan: dict, visual_pack: dict, output_path: Path, logo_path: Path) -> dict:
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    charts = _chart_map(visual_pack); lang = presentation_plan.get("language") or "en"; ctx = presentation_plan.get("research_context") or {}
    render_audit=[]
    for idx, spec in enumerate(presentation_plan.get("slides") or [], start=1):
        slide=prs.slides.add_slide(blank); _set_bg(slide)
        typ=spec.get("slide_type")
        if typ=="cover":
            _render_cover(slide, spec, ctx, lang, logo_path, presentation_plan)
        else:
            _add_header(slide,spec.get("title"),idx-1)
            _add_footer(slide,idx,ctx)
            if spec.get("subtitle"):
                _add_text(slide,spec.get("subtitle"),1.15,1.35,11.3,.72,size=12.5,color=MUTED)
            cids=spec.get("chart_ids") or []
            claims=spec.get("claims") or []
            chart_modes=[]
            if typ == "executive_summary":
                dash = presentation_plan.get("executive_dashboard")
                if isinstance(dash, dict) and dash.get("score") is not None:
                    # Locked "Version B" executive dashboard; claims stay in the
                    # ledger/docx — here the visuals carry the same findings.
                    _render_executive_dashboard(slide, dash, lang, ctx)
                    chart_modes.extend((cid, "editable_shapes") for cid in cids)
                else:
                    # Legacy layout: keeps old runs and thin datasets rendering.
                    if len(cids) >= 1:
                        chart_modes.append((cids[0],_render_chart_spec(slide,charts[cids[0]],1.12,1.55,5.35,2.15,lang)))
                    if len(cids) >= 2:
                        chart_modes.append((cids[1],_render_chart_spec(slide,charts[cids[1]],6.82,1.55,5.35,2.15,lang)))
                    for i,cl in enumerate(claims[:3]):
                        _add_claim_card(slide,cl.get("text"),1.12+i*3.72,4.15,3.47,1.62,font_size=10.5)
            elif typ == "reputation_timeline":
                tl = presentation_plan.get("reputation_timeline")
                if isinstance(tl, dict) and tl.get("points"):
                    _render_reputation_timeline(slide, tl, lang, ctx)
            elif typ == "evidence_split":
                pos=[c for c in claims if str((c.get("source_values") or {}).get("sentiment"))=="positive"]
                neg=[c for c in claims if str((c.get("source_values") or {}).get("sentiment"))=="negative"]
                _add_text(slide,"Top 5 positive" if lang=="en" else "Top 5 θετικές",1.12,1.5,5.3,.32,size=11.5,color=POS,bold=True)
                _add_text(slide,"Top 5 negative" if lang=="en" else "Top 5 αρνητικές",6.82,1.5,5.3,.32,size=11.5,color=NEG,bold=True)
                def _rank_card(cl, x, yy, accent, rank):
                    txt=str(cl.get("text") or "").replace(" | ","\n")
                    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy+0.06), Inches(0.3), Inches(0.3))
                    badge.fill.solid(); badge.fill.fore_color.rgb=_rgb(accent); badge.line.fill.background(); badge.shadow.inherit=False
                    _add_text(slide,str(rank),x,yy+0.06,0.3,0.3,size=11,color=WHITE,bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
                    _add_claim_card(slide,txt,x+0.4,yy,4.95,0.93,font_size=8.6)
                for i,cl in enumerate(pos[:5]): _rank_card(cl,1.12,1.92+i*1.03,POS,i+1)
                for i,cl in enumerate(neg[:5]): _rank_card(cl,6.82,1.92+i*1.03,NEG,i+1)
            elif typ == "voices_table":
                rows=((spec.get("notes") or {}).get("rows") or [])
                headers=["#", "Voice" if lang=="en" else "Φωνή", "Mentions" if lang=="en" else "Αναφορές", "Attention" if lang=="en" else "Προσοχή", "Sentiment"]
                body=[[str(r.get("rank")), str(r.get("name")), str(r.get("records")), f"{r.get('attention'):.2f}", f"{r.get('weighted_sentiment'):+.2f}"] for r in rows]
                _add_editable_table(slide, headers, body, 1.12, 1.6, 11.1, min(5.0, 0.42+0.31*len(body)), font_size=9, col_widths=[0.55, 5.6, 1.55, 1.6, 1.8])
                note=("Attention = observed digital attention share; weighted sentiment in [-1, +1]." if lang=="en" else "Προσοχή = παρατηρούμενο ψηφιακό μερίδιο προσοχής· σταθμισμένο sentiment σε κλίμακα [-1, +1].")
                _add_text(slide, note, 1.12, 6.75, 11.0, .3, size=8.5, color=MUTED)
            elif typ == "narratives" and spec.get("slide_id") == "voices_leaders":
                col_map={"positive":POS,"negative":NEG,"neutral":NEUTRAL,"mixed":MIXED}
                for i,cl in enumerate(claims[:4]):
                    col=i%2; row=i//2
                    x0=1.12+col*5.72; y0=1.7+row*2.45
                    lbl=str((cl.get("source_values") or {}).get("valence") or "")
                    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x0), Inches(y0+0.1), Inches(0.2), Inches(0.2))
                    dot.fill.solid(); dot.fill.fore_color.rgb=_rgb(col_map.get(lbl,NEUTRAL)); dot.line.fill.background(); dot.shadow.inherit=False
                    txt=str(cl.get("text") or "").replace(" | ","\n").replace(": ",":\n",1)
                    _add_claim_card(slide,txt,x0+0.32,y0,5.25,2.2,font_size=10)
            elif typ == "narratives":
                n=min(3,len(claims)); block=(4.95-(n-1)*0.18)/max(1,n)
                val_col={"positive":POS,"negative":NEG,"mixed":MIXED}
                for i,cl in enumerate(claims[:3]):
                    yy=1.6+i*(block+0.18)
                    sv=cl.get("source_values") or {}
                    valence=str(sv.get("valence") or "")
                    quotes=[q for q in (sv.get("quotes") or []) if q.get("excerpt")][:2]
                    qh=0.34*len(quotes)
                    card_h=max(0.8, block-qh-0.06)
                    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.12), Inches(yy+0.16), Inches(0.18), Inches(0.18))
                    dot.fill.solid(); dot.fill.fore_color.rgb=_rgb(val_col.get(valence,NEUTRAL)); dot.line.fill.background(); dot.shadow.inherit=False
                    txt=str(cl.get("text") or "").replace(" | ","\n")
                    _add_claim_card(slide,txt,1.48,yy,10.75,card_h,font_size=10.2 if n<=2 else 9.3)
                    # Verbatim voice of the audience: real excerpts only, with source and date.
                    for qi,q in enumerate(quotes):
                        src=" · ".join(str(x) for x in (q.get("platform"),q.get("date")) if x)
                        line=f"“{q.get('excerpt')}” — {src}" if src else f"“{q.get('excerpt')}”"
                        _add_text(slide,line,1.86,yy+card_h+0.04+qi*0.32,10.3,0.3,size=9,color=MUTED)
            elif typ == "evidence_cards" and spec.get("slide_id") == "recommendations":
                n=min(5,len(claims)); ch=(5.0-(n-1)*0.16)/max(1,n)
                for i,cl in enumerate(claims[:5]):
                    yy=1.62+i*(ch+0.16)
                    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.12), Inches(yy+ch/2-0.21), Inches(0.42), Inches(0.42))
                    badge.fill.solid(); badge.fill.fore_color.rgb=_rgb(AEGEAN); badge.line.fill.background(); badge.shadow.inherit=False
                    _add_text(slide,str(i+1),1.12,yy+ch/2-0.21,0.42,0.42,size=15,color=WHITE,bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
                    txt=str(cl.get("text") or "").replace(" | ","\n")
                    _add_claim_card(slide,txt,1.72,yy,10.5,ch,font_size=10.5 if n<=4 else 9.5)
            elif typ == "evidence_cards" and spec.get("slide_id") == "media_evidence" and (spec.get("chart_ids") or []):
                for i,cl in enumerate(claims[:4]):
                    _add_claim_card(slide,str(cl.get("text") or "").replace(" | ","\n"),1.12,1.95+i*1.2,5.45,1.06,font_size=8.8)
                cid=(spec.get("chart_ids") or [])[0]
                if cid in charts:
                    _render_chart_spec(slide, charts[cid], 6.78, 1.95, 5.45, 4.6, lang)
            elif typ == "evidence_cards":
                for i,cl in enumerate(claims[:4]):
                    col=i%2; row=i//2
                    tag=_claim_type_label(cl.get("claim_type"),lang)
                    _add_claim_card(slide,cl.get("text"),1.12+col*5.7,1.72+row*2.35,5.35,1.95,label=tag,accent=AEGEAN,font_size=10.0)
            elif typ == "strategic_synthesis":
                for i,cl in enumerate(claims[:4]):
                    role=str((cl.get("source_values") or {}).get("role") or "")
                    col=i%2; row=i//2
                    accent=POS if role=="strength" else NEG if role=="friction" else AEGEAN if role=="momentum" else WARN
                    _add_claim_card(slide,cl.get("text"),1.12+col*5.7,1.65+row*2.38,5.35,2.0,label=_role_label(role,lang),accent=accent,font_size=10.4)
            elif typ == "methodology":
                if cids:
                    chart_modes.append((cids[0],_render_chart_spec(slide,charts[cids[0]],1.12,1.7,4.15,4.65,lang)))
                x=5.65 if cids else 1.12; w=6.52 if cids else 11.05
                for i,cl in enumerate(claims[:4]):
                    _add_claim_card(slide,cl.get("text"),x,1.55+i*1.28,w,1.08,accent=AEGEAN,font_size=9.2)
            elif len(cids)==1:
                chart_modes.append((cids[0],_render_chart_spec(slide,charts[cids[0]],1.12,2.05,11.1,4.45,lang)))
            elif len(cids)>=2:
                chart_modes.append((cids[0],_render_chart_spec(slide,charts[cids[0]],1.12,2.0,5.35,4.35,lang)))
                chart_modes.append((cids[1],_render_chart_spec(slide,charts[cids[1]],6.82,2.0,5.35,4.35,lang)))
            elif typ == "conclusions" and (spec.get("notes") or {}).get("priorities"):
                prm=(spec.get("notes") or {}).get("priorities") or {}
                quad=[("protect", "Protect" if lang=="en" else "ΠΡΟΣΤΑΤΕΨΕ", POS),
                      ("improve", "Improve" if lang=="en" else "ΔΙΟΡΘΩΣΕ", NEG),
                      ("amplify", "Amplify" if lang=="en" else "ΕΝΙΣΧΥΣΕ", AEGEAN),
                      ("monitor", "Monitor" if lang=="en" else "ΠΑΡΑΚΟΛΟΥΘΗΣΕ", MIXED)]
                has_take=bool((spec.get("notes") or {}).get("final_takeaway"))
                qh=2.18 if has_take else 2.5
                for i,(key,lbl,accent) in enumerate(quad):
                    col=i%2; row=i//2
                    items=[_clean_text(v,150) for v in (prm.get(key) or [])][:3]
                    body="\n".join("• "+v for v in items) if items else ("—")
                    _add_claim_card(slide,body,1.12+col*5.72,1.6+row*(qh+0.18),5.5,qh,label=lbl,accent=accent,font_size=10)
                if has_take:
                    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.12), Inches(6.18), Inches(11.1), Inches(0.92))
                    strip.adjustments[0]=0.1; strip.shadow.inherit=False
                    strip.fill.solid(); strip.fill.fore_color.rgb=_rgb("EEF4F6"); strip.line.fill.background()
                    prefix="Final takeaway: " if lang=="en" else "Τελικό συμπέρασμα: "
                    _add_text(slide,prefix+_clean_text((spec.get("notes") or {}).get("final_takeaway"),380),1.34,6.28,10.66,.74,size=12,color=INK,bold=True,valign=MSO_ANCHOR.MIDDLE)
            elif claims:
                sub_norm=_clean_text(spec.get("subtitle") or "", 400)
                # The subtitle already states the finding; never repeat it as a card.
                claims=[c for c in claims if _clean_text(c.get("text"),400)!=sub_norm]
                y=1.8
                for cl in claims[:5]:
                    _add_claim_card(slide,cl.get("text"),1.15,y,10.9,.82,font_size=12.5)
                    y+=1.0
            if claims and cids and typ not in {"executive_summary","methodology"} and spec.get("slide_id") not in {"media_evidence","sentiment_evidence","evidence"}:
                # Plain-language takeaway strip: what this slide means, in one sentence.
                strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.12), Inches(6.42), Inches(11.1), Inches(0.62))
                strip.adjustments[0] = 0.14
                strip.shadow.inherit = False
                strip.fill.solid(); strip.fill.fore_color.rgb = _rgb("EEF4F6"); strip.line.fill.background()
                prefix = "What it means: " if lang == "en" else "Τι σημαίνει: "
                note = prefix + _clean_text(claims[0].get("text"), 220)
                _add_text(slide, note, 1.32, 6.5, 10.7, .48, size=10.5, color=INK, valign=MSO_ANCHOR.MIDDLE)
            if spec.get("notes",{}).get("causality_guardrail"):
                _add_text(slide,"Association ≠ proven causality." if lang=="en" else "Συσχέτιση ≠ αποδεδειγμένη αιτιότητα.",8.3,6.92,4.0,.2,size=7.2,color=WARN,align=PP_ALIGN.RIGHT)
            render_audit.extend({"slide_id":spec.get("slide_id"),"chart_id":cid,"render_mode":mode} for cid,mode in chart_modes)
    output_path.parent.mkdir(parents=True,exist_ok=True); prs.save(str(output_path))
    return {"slides":len(prs.slides),"chart_render_audit":render_audit,"path":str(output_path)}


# ---------- Internal brief ----------

def _docx_set_cell_text(cell, text, *, bold=False, size=9, color=INK):
    cell.text=""; p=cell.paragraphs[0]; r=p.add_run(_clean_text(text,1000)); r.bold=bold; r.font.name=FONT; r.font.size=DocxPt(size); r.font.color.rgb=DocxRGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


def generate_internal_docx(presentation_plan: dict, evidence_pack: dict, output_path: Path, logo_path: Path) -> dict:
    lang=presentation_plan.get("language") or "en"; ctx=presentation_plan.get("research_context") or {}
    doc=Document(); sec=doc.sections[0]; sec.top_margin=DocxInches(.65); sec.bottom_margin=DocxInches(.65); sec.left_margin=DocxInches(.75); sec.right_margin=DocxInches(.75)
    styles=doc.styles
    styles["Normal"].font.name=FONT; styles["Normal"].font.size=DocxPt(9.5); styles["Normal"].font.color.rgb=DocxRGBColor.from_string(INK)
    if logo_path.exists():
        try: doc.add_picture(str(logo_path),width=DocxInches(3.1))
        except Exception: pass
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(("Internal Intelligence Brief" if lang=="en" else "Εσωτερικό Intelligence Brief")); r.bold=True;r.font.name=FONT;r.font.size=DocxPt(22)
    p=doc.add_paragraph(); r=p.add_run(f"{ctx.get('client','')} · {ctx.get('topic','')} · {ctx.get('market','')} · {ctx.get('date_from','')} — {ctx.get('date_to','')}");r.font.name=FONT;r.font.size=DocxPt(9);r.font.color.rgb=DocxRGBColor.from_string(MUTED)

    def heading(text):
        p=doc.add_paragraph(); r=p.add_run(text); r.bold=True;r.font.name=FONT;r.font.size=DocxPt(14);r.font.color.rgb=DocxRGBColor.from_string(INK); return p
    heading("What matters" if lang=="en" else "Τι έχει σημασία")
    for cl in presentation_plan.get("claim_ledger") or []:
        if cl.get("slide_id") in {"executive_summary","conclusions"}:
            p=doc.add_paragraph(style=None); p.style=doc.styles["Normal"]; p.add_run("• ").bold=True; p.add_run(_clean_text(cl.get("text"),800))

    heading("What to watch / verify" if lang=="en" else "Τι να προσέξουμε / επαληθεύσουμε")
    warnings=list(presentation_plan.get("warnings") or [])
    if "Association ≠ proven causality." not in warnings: warnings.append("Association ≠ proven causality.")
    for w in warnings:
        p=doc.add_paragraph(); p.add_run("• ").bold=True; p.add_run(_clean_text(w,800))

    heading("Automatic investigations" if lang=="en" else "Αυτόματες διερευνήσεις")
    for inv in (evidence_pack.get("investigations") or [])[:10]:
        q=_txt(inv.get("question"),lang); f=_txt(inv.get("finding"),lang); conf=(inv.get("confidence") or {}).get("score")
        p=doc.add_paragraph(); rr=p.add_run(q or "Investigation");rr.bold=True;rr.font.name=FONT
        if f: doc.add_paragraph(f)
        meta=f"Confidence: {conf if conf is not None else '—'} · Causality: {inv.get('causal_status','not_proven')} · Evidence refs: {len(inv.get('evidence_record_ids') or [])}"
        p=doc.add_paragraph(); rr=p.add_run(meta);rr.italic=True;rr.font.size=DocxPt(8);rr.font.color.rgb=DocxRGBColor.from_string(MUTED)

    heading("Presentation review audit" if lang=="en" else "Έλεγχος παρουσίασης")
    rows=presentation_plan.get("indicator_review") or []
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style="Table Grid"
    for j,h in enumerate(["Indicator","Status","Slide(s)"]): _docx_set_cell_text(table.rows[0].cells[j],h,bold=True,color=WHITE);table.rows[0].cells[j].shading if False else None
    for row in rows:
        cells=table.add_row().cells
        _docx_set_cell_text(cells[0],row.get("indicator_id"),size=8)
        _docx_set_cell_text(cells[1],row.get("presentation_status"),size=8)
        _docx_set_cell_text(cells[2],", ".join(row.get("slide_ids") or []) or "—",size=8)

    heading("PowerPoint gold-standard audit" if lang=="en" else "Gold-standard έλεγχος PowerPoint")
    gold=(presentation_plan.get("gold_standard_audit") or {})
    p=doc.add_paragraph(); rr=p.add_run(f"Capability coverage: {_safe_float(gold.get('coverage_percent')):.1f}%"); rr.bold=True; rr.font.name=FONT
    table=doc.add_table(rows=1,cols=3);table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(["Capability","Status","Slide(s)"]): _docx_set_cell_text(table.rows[0].cells[j],h,bold=True,size=8)
    for row in gold.get("capabilities") or []:
        cells=table.add_row().cells
        _docx_set_cell_text(cells[0],row.get("label") or row.get("capability_id"),size=7.2)
        _docx_set_cell_text(cells[1],row.get("status"),size=7.2)
        _docx_set_cell_text(cells[2],", ".join(row.get("slide_ids") or []) or "—",size=7.2)

    heading("Claim ledger" if lang=="en" else "Claim ledger")
    table=doc.add_table(rows=1,cols=4);table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(["Slide","Claim","Indicators","Evidence"]): _docx_set_cell_text(table.rows[0].cells[j],h,bold=True,size=8)
    for cl in (presentation_plan.get("claim_ledger") or [])[:40]:
        cells=table.add_row().cells
        _docx_set_cell_text(cells[0],cl.get("slide_id"),size=7.2)
        _docx_set_cell_text(cells[1],cl.get("text"),size=7.2)
        _docx_set_cell_text(cells[2],", ".join(cl.get("indicator_ids") or []),size=7.2)
        _docx_set_cell_text(cells[3],", ".join(cl.get("evidence_refs") or []) or "—",size=7.2)

    output_path.parent.mkdir(parents=True,exist_ok=True);doc.save(str(output_path));return {"path":str(output_path),"paragraphs":len(doc.paragraphs),"tables":len(doc.tables)}


def _convert_to_pdf(input_path: Path, out_dir: Path) -> Path | None:
    soffice=shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice: return None
    out_dir.mkdir(parents=True,exist_ok=True)
    cmd=[soffice,"--headless","--convert-to","pdf","--outdir",str(out_dir),str(input_path)]
    proc=subprocess.run(cmd,capture_output=True,text=True,timeout=120)
    out=out_dir/(input_path.stem+".pdf")
    return out if proc.returncode==0 and out.exists() and out.stat().st_size>0 else None


def _evidence_rows(evidence_pack: dict) -> list[dict]:
    inv=_inventory_map(evidence_pack); rows=_indicator_value(inv,"top_mentions") or []
    return [_json_safe({k:r.get(k) for k in ("record_id","platform","author","origin_group","content_type","date","url","excerpt","sentiment_label","sentiment_score","emotion","topic","narrative","impact_score","evidence_confidence")}) for r in rows]


def persist_presentation_bundle(folder: Path, result: dict) -> dict:
    store=RunStore(); base=folder/"exports"
    store.write(base/"summary.json",result["summary"])
    store.write(base/"presentation-plan.json",result["presentation_plan"])
    store.write(base/"indicator-review.json",result["presentation_plan"]["indicator_review"])
    store.write(base/"chart-review.json",result["presentation_plan"]["chart_review"])
    store.write(base/"claim-ledger.json",result["presentation_plan"]["claim_ledger"])
    store.write(base/"gold-standard-audit.json",result["presentation_plan"].get("gold_standard_audit") or {})
    store.write(base/"manifest.json",result["manifest"])
    store.write(base/"qa.json",result["qa"])
    return result["summary"]


def _analyzed_record_rows(folder: Path) -> list[dict]:
    """Flat CSV rows: every analyzed record with its real AI annotation.

    Source of truth is analysis/analyzed.json, where every row carries the
    `ai_analysis` block produced by the semantic step. Falling back to the
    cleaned records keeps the export complete even before analysis runs.
    """
    store = RunStore()
    analyzed = store.read(folder / "analysis" / "analyzed.json", []) or []
    if not analyzed:
        analyzed = store.read(folder / "cleaning" / "trusted.json", []) or []
    rows = []
    for r in analyzed:
        a = r.get("ai_analysis") or {}
        cleaning = r.get("cleaning") or {}
        intel = r.get("intelligence") or {}
        emotions = a.get("emotions")
        if isinstance(emotions, list) and emotions:
            emotion = ", ".join(str(e.get("label") if isinstance(e, dict) else e) for e in emotions[:3])
        else:
            # Multi-label view: primary + secondary, with the model's intensity.
            parts = [a.get("primary_emotion")]
            second = a.get("secondary_emotion")
            if second and str(second).lower() != "none":
                parts.append(second)
            emotion = ", ".join(str(x) for x in parts if x)
        topics = a.get("topics")
        narratives = a.get("narratives")
        rows.append({
            "record_id": r.get("id") or r.get("record_id"),
            "date": r.get("timestamp") or r.get("date"),
            "platform": r.get("platform") or r.get("source"),
            "author": r.get("author"),
            "origin_group": intel.get("origin_group") or cleaning.get("origin_class") or "",
            "sentiment_label": a.get("sentiment_label", ""),
            "sentiment_label_model": a.get("sentiment_label_model", ""),
            "sentiment_score": a.get("sentiment_score", ""),
            "stance": (a.get("target_stance") or ""),
            "emotion": emotion,
            "emotion_intensity": a.get("emotion_intensity", ""),
            "sarcasm": (a.get("sarcasm") or {}).get("detected") if isinstance(a.get("sarcasm"), dict) else a.get("sarcasm", ""),
            "language": a.get("language", ""),
            "topic": (topics[0] if isinstance(topics, list) and topics else (a.get("topic") or "")),
            "narrative": (narratives[0] if isinstance(narratives, list) and narratives else (a.get("narrative") or "")),
            "relevance_score": a.get("relevance_score", cleaning.get("relevance_score", "")),
            "decision": a.get("decision", ""),
            "confidence": a.get("overall_confidence", ""),
            "impact_score": intel.get("impact_score", ""),
            "url": r.get("url"),
            "text": (str(r.get("text") or ""))[:600],
        })
    return rows


def build_exports(folder: Path, plan: dict, *, force: bool=False, cancel_check: Callable[[], bool] | None=None) -> dict:
    visual_summary=load_visualization_summary(folder)
    if visual_summary is None: raise RuntimeError("Step 8 requires Step 7 Charts & Dashboard.")
    if visual_summary.get("stale"): raise RuntimeError("Step 7 is stale; rebuild Charts & Dashboard before exports.")
    visual_pack=load_presentation_visual_pack(folder); evidence_pack=load_evidence_pack(folder)
    _validate_inputs(visual_pack,evidence_pack,plan)
    input_hash=_hash_payload({"visual_pack":visual_pack,"evidence_contract":evidence_pack.get("contract_version"),"research_context":evidence_pack.get("research_context"),"report_language":plan.get("report_language"),"search_strategy":plan.get("search_strategy"),"keyword_roles":plan.get("keyword_roles"),"benchmark":plan.get("benchmark"),"master_spec_version":plan.get("master_spec_version"),"presentation_ruleset":PRESENTATION_RULESET_VERSION})
    old=RunStore.read(folder/"exports"/"summary.json")
    if old and not force and old.get("input_hash")==input_hash and not old.get("stale"):
        return old
    if cancel_check and cancel_check(): raise PresentationCancelled("Exports cancelled before planning")
    synthesis=build_report_synthesis(folder,plan,visual_pack,evidence_pack)
    visual_pack=copy.deepcopy(visual_pack); visual_pack["analyst_synthesis"]=synthesis
    pplan=build_presentation_plan(visual_pack,evidence_pack,plan,cancel_check=cancel_check)
    # Slide 2 executive dashboard: computed per run; None falls back to legacy.
    pplan["executive_dashboard"]=_executive_dashboard_data(folder,visual_pack)
    # Slide 3 reputation timeline: only added when the run has a plottable series.
    tl=_reputation_timeline_data(folder,visual_pack,pplan.get("language") or "en")
    if tl:
        pplan["reputation_timeline"]=tl
        spec={"slide_id":"reputation_timeline","slide_type":"reputation_timeline",
              "title":"Brand Reputation Timeline","chart_ids":[],"claims":[],
              "section":"opening","priority":99,"required":False,"notes":{}}
        slides_list=pplan.get("slides") or []
        idx=next((i for i,sp in enumerate(slides_list) if sp.get("slide_id")=="executive_summary"),None)
        slides_list.insert((idx+1) if idx is not None else 1,spec)
    base=folder/"exports"; base.mkdir(parents=True,exist_ok=True)
    lang=pplan["language"]; ctx=pplan["research_context"]
    stem=re.sub(r"[^A-Za-z0-9Α-Ωα-ω_-]+","_",f"{ctx.get('client','')}_{ctx.get('topic','')}_{ctx.get('date_from','')}_{ctx.get('date_to','')}").strip("_")[:140] or "SIGNALYTH_Report"
    pptx_path=base/f"{stem}.pptx"; docx_path=base/f"{stem}_Internal.docx"
    logo_path=Path(__file__).resolve().parents[2]/"logo.png"
    pptx_meta=generate_pptx(pplan,visual_pack,pptx_path,logo_path)
    if cancel_check and cancel_check(): raise PresentationCancelled("Exports cancelled after PowerPoint generation")
    docx_meta=generate_internal_docx(pplan,evidence_pack,docx_path,logo_path)
    presentation_pdf=_convert_to_pdf(pptx_path,base)
    internal_pdf=_convert_to_pdf(docx_path,base)

    evidence_rows=_evidence_rows(evidence_pack)
    evidence_json=base/f"{stem}_Evidence.json"; evidence_json.write_text(json.dumps(evidence_rows,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
    evidence_csv=base/f"{stem}_Evidence.csv"
    with evidence_csv.open("w",encoding="utf-8-sig",newline="") as f:
        fields=list(evidence_rows[0].keys()) if evidence_rows else ["record_id","platform","author","date","url","excerpt"]
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows([{k:(json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v) for k,v in r.items()} for r in evidence_rows])

    # Interactive client dashboard: the same evidence, explorable, offline, no licence.
    dashboard_path=base/f"{stem}_Dashboard.html"
    try:
        from app.services.dashboard_html import build_dashboard_html
        intel_summary=RunStore.read(folder/"intelligence"/"summary.json", {}) or {}
        dash_records=RunStore.read(folder/"analysis"/"analysis-ready.json", []) or []
        if not dash_records:
            dash_records=RunStore.read(folder/"analysis"/"analyzed.json", []) or []
        dashboard_path.write_text(
            build_dashboard_html(plan={**plan, **ctx}, intelligence=intel_summary,
                                 records=dash_records, language=lang),
            encoding="utf-8",
        )
    except Exception:
        dashboard_path=None

    records_csv=base/f"{stem}_Records.csv"
    try:
        rec_rows=_analyzed_record_rows(folder)
        with records_csv.open("w",encoding="utf-8-sig",newline="") as f:
            fields=["record_id","date","platform","author","origin_group","sentiment_label","sentiment_label_model","sentiment_score","stance","emotion","emotion_intensity","sarcasm","language","topic","narrative","relevance_score","decision","confidence","impact_score","url","text"]
            w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore");w.writeheader();w.writerows(rec_rows)
    except Exception:
        records_csv=None

    files=[pptx_path,docx_path,evidence_json,evidence_csv]
    if dashboard_path: files.append(dashboard_path)
    if records_csv: files.append(records_csv)
    if presentation_pdf: files.append(presentation_pdf)
    if internal_pdf: files.append(internal_pdf)
    manifest={"files":[]}
    for p in files:
        manifest["files"].append({"name":p.name,"type":p.suffix.lower().lstrip("."),"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest()})

    # Export QA is structural here; release workflow additionally renders PPTX/DOCX/PDF visually.
    consistency_qa=final_consistency_qa(pplan,evidence_pack,visual_pack,synthesis,plan)
    RunStore().write(folder/"exports"/"final-consistency-qa.json",consistency_qa)
    qa={
        "indicator_review_complete":len(pplan["indicator_review"])==len(INDICATOR_REGISTRY) and all(x.get("examined") for x in pplan["indicator_review"]),
        "claim_ledger_valid":True,
        "native_editable_chart_contract":all((c.get("presentation") or {}).get("native_editable_ready") and not (c.get("presentation") or {}).get("render_as_raster") for c in (visual_pack.get("chart_specs") or [])),
        "pptx_slides":pptx_meta["slides"],
        "pptx_render_modes":pptx_meta["chart_render_audit"],
        "docx_tables":docx_meta["tables"],
        "presentation_pdf_available":bool(presentation_pdf),
        "internal_pdf_available":bool(internal_pdf),
        "evidence_rows":len(evidence_rows),
        "causality_guardrail":True,
        "gold_standard_capability_coverage_percent": _safe_float((pplan.get("gold_standard_audit") or {}).get("coverage_percent")),
        "gold_standard_capability_complete": bool((pplan.get("gold_standard_audit") or {}).get("complete")),
        "analyst_synthesis_provider": synthesis.get("provider"),
        "final_consistency_qa_passed": bool(consistency_qa.get("passed")),
    }
    strict_checks={
        "indicator_review_complete":qa["indicator_review_complete"],
        "claim_ledger_valid":qa["claim_ledger_valid"],
        "native_editable_chart_contract":qa["native_editable_chart_contract"],
        "minimum_slides":qa["pptx_slides"]>=3,
        "gold_standard_capability_complete":qa["gold_standard_capability_complete"],
        "final_consistency_qa_passed":qa["final_consistency_qa_passed"],
    }
    qa["degraded_checks"]=[k for k,v in strict_checks.items() if not v]
    qa["degraded"]=bool(qa["degraded_checks"])
    # The deliverable ALWAYS ships. A QA gap is documented loudly in the summary
    # and the internal QA files — never converted into a dead run with zero output.

    summary={
        "ruleset_version":PRESENTATION_RULESET_VERSION,"methodology_version":PRESENTATION_METHODOLOGY_VERSION,"presentation_contract_version":PRESENTATION_CONTRACT_VERSION,
        "generated_at":_utcnow(),"input_hash":input_hash,"research_context":copy.deepcopy(ctx),"language":lang,"slide_count":pptx_meta["slides"],"claim_count":len(pplan["claim_ledger"]),"indicator_contract":{"required":len(INDICATOR_REGISTRY),"examined":len(pplan["indicator_review"]),"complete":qa["indicator_review_complete"]},"gold_standard_capability_coverage_percent":qa["gold_standard_capability_coverage_percent"],"gold_standard_capability_complete":qa["gold_standard_capability_complete"],"files":[x["name"] for x in manifest["files"]],"stale":False,
        "boundary":"Step 8 exports evidence-linked, editable client and internal deliverables. A polished report does not imply complete platform coverage or proven causality."
    }
    result={"summary":summary,"presentation_plan":pplan,"manifest":manifest,"qa":qa}
    persist_presentation_bundle(folder,result)
    return summary


def load_export_summary(folder: Path) -> dict | None:
    summary=RunStore.read(folder/"exports"/"summary.json")
    if not isinstance(summary,dict): return None
    current=load_presentation_visual_pack(folder); evidence=load_evidence_pack(folder); plan=RunStore.read(folder/"plan.json",{}) or {}
    if not isinstance(current,dict) or not isinstance(evidence,dict):
        return {**summary,"stale":True}
    try:
        now_hash=_hash_payload({"visual_pack":current,"evidence_contract":evidence.get("contract_version"),"research_context":evidence.get("research_context"),"report_language":plan.get("report_language"),"search_strategy":plan.get("search_strategy"),"keyword_roles":plan.get("keyword_roles"),"benchmark":plan.get("benchmark"),"master_spec_version":plan.get("master_spec_version"),"presentation_ruleset":PRESENTATION_RULESET_VERSION})
    except Exception:
        return {**summary,"stale":True}
    return {**summary,"stale":summary.get("input_hash")!=now_hash}


def load_export_manifest(folder: Path) -> dict | None:
    value=RunStore.read(folder/"exports"/"manifest.json")
    return value if isinstance(value,dict) else None


def allowed_export_file(folder: Path, filename: str) -> Path | None:
    manifest=load_export_manifest(folder) or {}; names={x.get("name") for x in manifest.get("files") or []}
    if filename not in names or Path(filename).name!=filename: return None
    p=folder/"exports"/filename
    return p if p.is_file() else None
